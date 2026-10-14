import os
import pandas as pd
import re
import datetime as dt
import shutil
import hashlib
import json
import tempfile
from pathlib import Path

from core.exporters import save_workbook_atomic
from core.persistence import write_json_atomic


SPANISH_MONTHS = {
    "ENERO": 1,
    "FEBRERO": 2,
    "MARZO": 3,
    "ABRIL": 4,
    "MAYO": 5,
    "JUNIO": 6,
    "JULIO": 7,
    "AGOSTO": 8,
    "SETIEMBRE": 9,
    "SEPTIEMBRE": 9,
    "OCTUBRE": 10,
    "NOVIEMBRE": 11,
    "DICIEMBRE": 12,
}


def normalize_document(value):
    """Normaliza DNI/CE sin convertir identificadores a flotantes."""
    if value is None or pd.isna(value):
        return ""
    text = str(value).strip()
    if not text or text.lower() in {"nan", "none", "null"}:
        return ""
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    digits = re.sub(r"\D", "", text)
    return digits.lstrip("0") or ("0" if digits else "")

def parse_date(date_str):
    if date_str is None or (not isinstance(date_str, (dt.date, dt.datetime)) and pd.isna(date_str)):
        return None, None

    if isinstance(date_str, (dt.datetime, dt.date)):
        return int(date_str.year), int(date_str.month)

    s = str(date_str).strip()
    if not s:
        return None, None

    # Si es solo una hora (ej. 17:18:06), no es fecha válida
    if re.match(r"^\d{1,2}:\d{2}(?::\d{2})?$", s):
        return None, None

    # Fechas ISO con hora, DD/MM/YYYY, DD-MM-YYYY y variantes de un dígito.
    match = re.search(r"(?<!\d)(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})", s)
    if match:
        day, month, year = map(int, match.groups())
        if 1 <= month <= 12 and 1 <= day <= 31:
            return year, month

    match = re.search(r"(?<!\d)(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", s)
    if match:
        year, month, day = map(int, match.groups())
        if 1 <= month <= 12 and 1 <= day <= 31:
            return year, month

    # DD/MM/YY (años de 2 dígitos, ej. 12/08/26)
    match = re.search(r"(?<!\d)(\d{1,2})[-/.](\d{1,2})[-/.](\d{2})\b", s)
    if match:
        day, month, yy = map(int, match.groups())
        if 1 <= month <= 12 and 1 <= day <= 31:
            year = 2000 + yy if 0 <= yy <= 40 else 1900 + yy
            return year, month

    month_names = "|".join(SPANISH_MONTHS)
    match = re.search(rf"(?i)\b(\d{{1,2}})\s+(?:DE\s+)?({month_names})\s+(?:(?:DE|DEL)\s+)?(\d{{4}})\b", s)
    if match:
        return int(match.group(3)), SPANISH_MONTHS[match.group(2).upper()]

    return None, None


def parse_full_date(date_value) -> dt.date | None:
    """Obtiene fecha exacta cuando está disponible; evita inventar el día."""
    if date_value is None or (not isinstance(date_value, (dt.date, dt.datetime)) and pd.isna(date_value)):
        return None
    if isinstance(date_value, dt.datetime):
        return date_value.date()
    if isinstance(date_value, dt.date):
        return date_value
    text = str(date_value).strip()
    match = re.search(r"(?<!\d)(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})", text)
    if match:
        day, month, year = map(int, match.groups())
    else:
        match = re.search(r"(?<!\d)(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", text)
        if not match:
            return None
        year, month, day = map(int, match.groups())
    try:
        return dt.date(year, month, day)
    except ValueError:
        return None


def _read_reference_table(reference_path: str | Path) -> pd.DataFrame:
    path = Path(reference_path)
    if path.suffix.lower() in {".xlsx", ".xls"}:
        return pd.read_excel(path)
    try:
        frame = pd.read_csv(path, sep=";", encoding="utf-8")
        if len(frame.columns) <= 1:
            frame = pd.read_csv(path, sep=",", encoding="utf-8")
        return frame
    except Exception:
        return pd.read_csv(path, sep=",", encoding="utf-8")


def validate_comparison_inputs(ocr_excel_path: str, reference_csv: str) -> dict:
    """Valida entradas de Hito 4 antes de modificar el Excel."""
    from openpyxl import load_workbook

    ocr_path = Path(ocr_excel_path)
    reference_path = Path(reference_csv)
    if ocr_path.suffix.lower() != ".xlsx":
        raise ValueError("El archivo de extracción debe ser un Excel .xlsx.")
    if not ocr_path.exists() or not reference_path.exists():
        raise ValueError("No existe uno de los archivos seleccionados.")

    reference = _read_reference_table(reference_csv)
    reference_columns = {str(column).strip() for column in reference.columns}
    has_reference_date = "Fecha de Muestra" in reference_columns or {"Anio", "Mes"}.issubset(reference_columns)
    if "Num Doc." not in reference_columns or not has_reference_date:
        raise ValueError("La referencia SAIRC debe incluir 'Num Doc.' y fecha de muestra o columnas 'Anio'/'Mes'.")

    workbook = load_workbook(ocr_path, read_only=True, data_only=False)
    try:
        if "SAIRC_Formulario2" not in workbook.sheetnames:
            raise ValueError("El Excel debe contener la hoja 'SAIRC_Formulario2'.")
        worksheet = workbook["SAIRC_Formulario2"]
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in next(worksheet.iter_rows(max_row=1))]
        required = {"documento_archivo_pdf", "fecha_toma_muestra"}
        missing = sorted(required - set(headers))
        if missing:
            raise ValueError(f"Faltan columnas obligatorias en el Excel: {', '.join(missing)}")
        comparable = sorted(ocr_key for ocr_key, ref_key in FIELD_MAP.items() if ref_key in reference_columns and ocr_key in headers)
        return {
            "filas_referencia": int(len(reference)),
            "filas_extraccion": max(0, int(worksheet.max_row) - 1) if worksheet.max_row is not None else sum(1 for _ in worksheet.iter_rows(min_row=2)),
            "campos_referencia": sorted(reference_columns & set(FIELD_MAP.values())),
            "columnas_ocr": headers,
            "campos_comparables": comparable,
        }
    finally:
        workbook.close()

import math

def safe_float(v):
    if pd.isna(v) or v is None:
        return None
    s = str(v).strip().lower()
    if s in ['nan', 'none', '', 'null']:
        return None
    try:
        f = float(s.replace(',', '.'))
        if math.isnan(f):
            return None
        return f
    except (ValueError, TypeError):
        return None


def numeric_values_equal(left, right, *, abs_tolerance=0.0001, rel_tolerance=0.000001):
    if left is None or right is None:
        return left is right
    return math.isclose(left, right, abs_tol=abs_tolerance, rel_tol=rel_tolerance)


FIELD_MAP = {
    'UREA_PRE': 'UPRE',
    'UREA_POST': 'UPOST',
    'HEMOGLOBINA': 'HB',
    'HEMATOCRITO': 'HTO',
    'CALCIO': 'CalcioSerico',
    'FOSFORO': 'FosforoSerico',
    'PTH': 'PTH',
    'ALBUMINA': 'AlbuminaSerica',
    'FOSFATASA_ALCALINA': 'FOSFATASA',
    'TGO_AST': 'TGO',
    'TGP_ALT': 'TGP',
    'HIERRO_SERICO': 'HIERROSERICO',
    'FERRITINA': 'FERRITINA'
}

GROUP1_KEYS = ['UREA_PRE', 'UREA_POST', 'HEMOGLOBINA', 'HEMATOCRITO']
PENALTY_AUDIT_FIELDS = {"HEMOGLOBINA", "UREA_PRE", "UREA_POST"}

LEGAL_ENTITY_SUFFIX = re.compile(
    r"(?:[\s,.-]+(?:S\.?A\.?C\.?|E\.?I\.?R\.?L\.?|S\.?R\.?L\.?|S\.?A\.?))+$",
    re.IGNORECASE,
)


def parse_ipress_from_relative_path(relative_path) -> tuple[str, str]:
    """Obtiene CUI y nombre de la IPRESS desde la primera carpeta de la ruta."""
    if relative_path is None or pd.isna(relative_path):
        return "", ""
    text = str(relative_path).strip().replace("\\", "/").strip("/")
    if not text:
        return "", ""
    parts = [part.strip() for part in text.split("/") if part.strip()]
    if len(parts) < 4:
        return "", ""
    first_folder = parts[0]
    match = re.match(r"^(\d+)\s*-\s*(.+?)\s*$", first_folder)
    if not match:
        return "", first_folder
    cui, name = match.groups()
    name = LEGAL_ENTITY_SUFFIX.sub("", name).strip(" ,-.")
    if name.isupper():
        name = name.title()
        for connector in (" De ", " Del ", " La ", " Las ", " Los ", " Y "):
            name = name.replace(connector, connector.lower())
    return cui, name


def normalize_ipress_code(value) -> str:
    """Normaliza CUI/RENAES conservándolo como identificador textual."""
    if value is None or pd.isna(value):
        return ""
    text = str(value).strip()
    if re.fullmatch(r"\d+\.0+", text):
        text = text.split(".", 1)[0]
    digits = re.sub(r"\D", "", text)
    return digits.lstrip("0") or ("0" if digits else "")


def resolve_source_ipress(relative_path, filename=None) -> tuple[str, str, str]:
    """Resuelve CUI por ruta y, para rutas históricas incompletas, por nombre del PDF."""
    cui, name = parse_ipress_from_relative_path(relative_path)
    normalized = normalize_ipress_code(cui)
    if normalized:
        return normalized, name, "RUTA_PDF"
    filename_text = str(filename or "").strip()
    match = re.search(r"_0*(\d{4,6})-\d{14}(?:\D|$)", filename_text)
    if match:
        return normalize_ipress_code(match.group(1)), name, "NOMBRE_PDF"
    return "", name, "NO_DISPONIBLE"


def _safe_relative_pdf_path(value) -> Path | None:
    if value is None or pd.isna(value):
        return None
    parts = [part for part in re.split(r"[\\/]+", str(value).strip()) if part not in {"", "."}]
    if not parts or any(part == ".." or ":" in part for part in parts):
        return None
    relative = Path(*parts)
    return relative if relative.suffix.lower() == ".pdf" else None


def discover_pdf_root(ocr_excel_path: str, relative_paths) -> Path | None:
    """Busca una raíz que contenga la primera ruta relativa, priorizando carpetas cercanas al resultado."""
    first_relative = next((_safe_relative_pdf_path(value) for value in relative_paths if _safe_relative_pdf_path(value)), None)
    if first_relative is None:
        return None
    excel_path = Path(ocr_excel_path).resolve()
    for parent in excel_path.parents:
        if (parent / first_relative).is_file():
            return parent
        if parent == parent.parent:
            break
        try:
            for child in parent.iterdir():
                if child.is_dir() and (child / first_relative).is_file():
                    return child
        except (OSError, PermissionError):
            continue
    return None


def _sanitize_path_for_filesystem(relative: Path) -> Path:
    """Sanitiza cada componente de la ruta relativa para evitar errores de Windows

    como [WinError 3] provocados por espacios o puntos al final de carpetas.
    """
    clean_parts = []
    for idx, part in enumerate(relative.parts):
        s = part.strip()
        if idx < len(relative.parts) - 1:
            s = s.rstrip(". ")
        else:
            s = s.rstrip(" ")
        clean_parts.append(s if s else part)
    return Path(*clean_parts)


def _to_win_long_path(path: Path | str) -> str:
    path_str = str(Path(path).resolve())
    if os.name == "nt" and not path_str.startswith("\\\\?\\"):
        if path_str.startswith("\\\\"):
            return "\\\\?\\UNC\\" + path_str[2:]
        return "\\\\?\\" + path_str
    return path_str


def _safe_mkdir(path: Path) -> None:
    try:
        path.mkdir(parents=True, exist_ok=True)
    except OSError:
        if os.name == "nt":
            os.makedirs(_to_win_long_path(path), exist_ok=True)
        else:
            raise


def _safe_exists(path: Path | str) -> bool:
    try:
        p = Path(path)
        if p.exists():
            return True
    except OSError:
        pass
    if os.name == "nt":
        try:
            return os.path.exists(_to_win_long_path(path))
        except OSError:
            return False
    return False


def _safe_is_file(path: Path | str) -> bool:
    try:
        p = Path(path)
        if p.is_file():
            return True
    except OSError:
        pass
    if os.name == "nt":
        try:
            return os.path.isfile(_to_win_long_path(path))
        except OSError:
            return False
    return False


def _safe_file_size(path: Path | str) -> int:
    try:
        return Path(path).stat().st_size
    except OSError:
        if os.name == "nt":
            return os.path.getsize(_to_win_long_path(path))
        raise


def _safe_copy2(source: Path, destination: Path) -> None:
    with tempfile.NamedTemporaryFile(dir=destination.parent, suffix=".pdf", delete=False) as stream:
        temporary = Path(stream.name)
    try:
        try:
            shutil.copy2(source, temporary)
            os.replace(temporary, destination)
        except OSError:
            if os.name != "nt":
                raise
            shutil.copy2(_to_win_long_path(source), _to_win_long_path(temporary))
            os.replace(_to_win_long_path(temporary), _to_win_long_path(destination))
    finally:
        temporary.unlink(missing_ok=True)


def _safe_move(source: Path, destination: Path) -> None:
    try:
        shutil.move(str(source), str(destination))
    except OSError:
        if os.name == "nt":
            shutil.move(_to_win_long_path(source), _to_win_long_path(destination))
        else:
            raise


def prepare_review_pdfs(
    findings: list[dict],
    pdf_root: str | Path,
    review_dir: str | Path,
    *,
    operation: str = "copy",
) -> dict:
    """Concentra PDF únicos con hallazgos, preservando la estructura y la trazabilidad."""
    if operation not in {"copy", "move"}:
        raise ValueError("La operación de revisión debe ser 'copy' o 'move'.")
    source_root = Path(pdf_root).resolve()
    destination_root = Path(review_dir).resolve()
    if not source_root.is_dir():
        raise ValueError(f"No existe la carpeta raíz de PDF: {source_root}")
    if source_root == destination_root or source_root in destination_root.parents or destination_root in source_root.parents:
        raise ValueError("Las carpetas fuente y de revisión deben estar separadas, sin contenerse entre sí.")

    manifest_path = destination_root / "copias_revision.json"
    previous_copies = {}
    if manifest_path.exists():
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        if manifest.get("source_root") != str(source_root):
            raise ValueError("La carpeta de revisión pertenece a otra raíz de PDF; selecciona una salida independiente.")
        previous_copies = manifest.get("copies", {})
        if not isinstance(previous_copies, dict):
            raise ValueError("El registro de copias de revisión no es válido.")
    managed_copies = dict(previous_copies)

    def pdf_digest(path: Path) -> str:
        with path.open("rb") as stream:
            return hashlib.file_digest(stream, "sha256").hexdigest()

    candidates: dict[Path, list[dict]] = {}
    for finding in findings:
        relative = _safe_relative_pdf_path(finding.get("Ruta_relativa"))
        if relative is not None:
            candidates.setdefault(relative, []).append(finding)

    stats = {
        "seleccionados": len(candidates),
        "copiados": 0,
        "movidos": 0,
        "ya_preparados": 0,
        "no_encontrados": 0,
        "rutas_invalidas": 0,
        "rutas_recuperadas": 0,
        "rutas_ambiguas": 0,
        "obsoletos_eliminados": 0,
    }
    try:
        ipress_directories = [path for path in source_root.iterdir() if path.is_dir()]
    except (OSError, PermissionError):
        ipress_directories = []
    expected_destinations: set[Path] = set()
    for relative, related_findings in candidates.items():
        source = (source_root / relative).resolve()
        resolved_relative = relative
        if not _safe_is_file(source):
            alt_relative = _sanitize_path_for_filesystem(relative)
            if alt_relative != relative and _safe_is_file(source_root / alt_relative):
                source = (source_root / alt_relative).resolve()
                resolved_relative = alt_relative
            else:
                recovered = [(directory / relative).resolve() for directory in ipress_directories if _safe_is_file(directory / relative)]
                if not recovered and alt_relative != relative:
                    recovered = [(directory / alt_relative).resolve() for directory in ipress_directories if _safe_is_file(directory / alt_relative)]
                if len(recovered) == 1:
                    source = recovered[0]
                    resolved_relative = source.relative_to(source_root)
                    stats["rutas_recuperadas"] += 1
                elif len(recovered) > 1:
                    stats["rutas_ambiguas"] += 1
                    for finding in related_findings:
                        finding["Estado_PDF_Revision"] = "RUTA_AMBIGUA"
                    continue
        destination_relative = _sanitize_path_for_filesystem(resolved_relative)
        destination = (destination_root / destination_relative).resolve()
        produced_copy = False
        try:
            source.relative_to(source_root)
            destination.relative_to(destination_root)
        except ValueError:
            stats["rutas_invalidas"] += 1
            status = "RUTA_INVALIDA"
        else:
            expected_destinations.add(destination)
            _safe_mkdir(destination.parent)
            relative_key = destination.relative_to(destination_root).as_posix()
            source_exists = _safe_is_file(source)
            destination_exists = _safe_is_file(destination)
            destination_digest = pdf_digest(destination) if destination_exists else None
            same_content = source_exists and destination_exists and pdf_digest(source) == destination_digest
            if source_exists and destination_exists and not same_content:
                if previous_copies.get(relative_key) != destination_digest:
                    raise FileExistsError(f"No se sobrescribirá un PDF ajeno o modificado en {destination}")
            if destination_exists and not source_exists:
                stats["ya_preparados"] += 1
                status = "DISPONIBLE_REVISION"
            elif not source_exists:
                stats["no_encontrados"] += 1
                status = "NO_ENCONTRADO"
            elif operation == "move":
                if _safe_exists(destination):
                    if not same_content:
                        raise FileExistsError(f"No se moverá el PDF: ya existe una copia diferente en {destination}")
                    stats["ya_preparados"] += 1
                    status = "DISPONIBLE_REVISION"
                else:
                    _safe_move(source, destination)
                    produced_copy = True
                    stats["movidos"] += 1
                    status = "DISPONIBLE_REVISION"
            else:
                if same_content:
                    stats["ya_preparados"] += 1
                    status = "DISPONIBLE_REVISION"
                else:
                    _safe_copy2(source, destination)
                    produced_copy = True
                    stats["copiados"] += 1
                    status = "DISPONIBLE_REVISION"
            if status == "DISPONIBLE_REVISION" and produced_copy:
                managed_copies[relative_key] = pdf_digest(destination)

        for finding in related_findings:
            finding["Ruta_PDF_Revision"] = str(destination) if status != "RUTA_INVALIDA" else ""
            finding["Estado_PDF_Revision"] = status
            cui_ipress, ipress_name = parse_ipress_from_relative_path(resolved_relative)
            if cui_ipress and ipress_name:
                finding["CUI_IPRESS"] = cui_ipress
                finding["Establecimiento"] = ipress_name
                if resolved_relative != relative:
                    finding["Origen_Establecimiento"] = "RUTA_PDF_RECUPERADA"

    # Persistir la propiedad antes de retirar copias obsoletas.
    write_json_atomic(manifest_path, {"source_root": str(source_root), "copies": managed_copies})
    if operation == "copy":
        for relative_key, digest in previous_copies.items():
            relative = _safe_relative_pdf_path(relative_key)
            if relative is None:
                continue
            existing = (destination_root / relative).resolve()
            if not existing.is_relative_to(destination_root) or existing in expected_destinations:
                continue
            if existing.is_file() and pdf_digest(existing) == digest:
                existing.unlink()
                stats["obsoletos_eliminados"] += 1
            managed_copies.pop(relative_key, None)
    write_json_atomic(manifest_path, {"source_root": str(source_root), "copies": managed_copies})
    return stats


def discrepancy_magnitude(field: str, ocr_value: float, sairc_value: float) -> tuple[float, float | None, str]:
    """Clasifica magnitud para priorización de auditoría; no representa un rango clínico."""
    absolute = abs(ocr_value - sairc_value)
    percentage = (absolute / abs(sairc_value) * 100) if sairc_value else None

    if field == "HEMOGLOBINA":
        if absolute >= 2:
            magnitude = "CRITICA"
        elif absolute >= 1:
            magnitude = "ALTA"
        elif absolute >= 0.5:
            magnitude = "MEDIA"
        else:
            magnitude = "BAJA"
    elif field == "UREA_PRE":
        if absolute >= 50 or (percentage is not None and percentage >= 30):
            magnitude = "CRITICA"
        elif absolute >= 20 or (percentage is not None and percentage >= 15):
            magnitude = "ALTA"
        elif absolute >= 10 or (percentage is not None and percentage >= 8):
            magnitude = "MEDIA"
        else:
            magnitude = "BAJA"
    elif field == "UREA_POST":
        if absolute >= 20 or (percentage is not None and percentage >= 40):
            magnitude = "CRITICA"
        elif absolute >= 10 or (percentage is not None and percentage >= 20):
            magnitude = "ALTA"
        elif absolute >= 5 or (percentage is not None and percentage >= 10):
            magnitude = "MEDIA"
        else:
            magnitude = "BAJA"
    elif percentage is None:
        if absolute >= 50:
            magnitude = "CRITICA"
        elif absolute >= 10:
            magnitude = "ALTA"
        elif absolute >= 1:
            magnitude = "MEDIA"
        else:
            magnitude = "BAJA"
    elif percentage >= 50:
        magnitude = "CRITICA"
    elif percentage >= 25:
        magnitude = "ALTA"
    elif percentage >= 10:
        magnitude = "MEDIA"
    else:
        magnitude = "BAJA"
    return round(absolute, 4), round(percentage, 2) if percentage is not None else None, magnitude


def _finding_priority(field: str | None, magnitude: str, finding_type: str) -> tuple[str, int, str]:
    if finding_type == "PDF_SIN_REGISTRO_SAIRC":
        return "CRITICA", 100, "PDF con resultados sin registro SAIRC correlacionado; posible riesgo de omisión sujeto a verificación."
    if finding_type == "REGISTRO_SAIRC_SIN_RESULTADOS":
        return "CRITICA", 98, "Existe la fila SAIRC, pero no contiene ninguno de los resultados comparables presentes en el PDF."
    if finding_type == "SAIRC_SIN_PDF":
        return "ALTA", 90, "Registro SAIRC sin informe PDF correlacionado; verificar archivo, documento y periodo."
    if finding_type in {"SIN_DOCUMENTO", "FECHA_INVALIDA", "CUI_IPRESS_NO_IDENTIFICADO"}:
        return "ALTA", 85, "La identidad, fecha o CUI/IPRESS impide correlacionar el PDF con SAIRC."
    if finding_type == "REFERENCIA_DUPLICADA":
        return "MEDIA", 65, "La misma llave documento-periodo aparece más de una vez en la referencia SAIRC."
    if finding_type in {"SAIRC_CAMPO_SIN_VALOR_PDF", "PDF_CAMPO_SIN_REGISTRO_SAIRC"}:
        if field in PENALTY_AUDIT_FIELDS:
            return "CRITICA", 95, "Falta un valor penalizable de hemoglobina o urea en uno de los dos orígenes."
        return "MEDIA", 55, "El examen tiene valor únicamente en uno de los dos orígenes."

    magnitude_score = {"CRITICA": 95, "ALTA": 80, "MEDIA": 60, "BAJA": 30}[magnitude]
    if field in PENALTY_AUDIT_FIELDS:
        return magnitude, magnitude_score, "Examen asociado a penalidad; priorizado por diferencia absoluta y porcentual."
    reduced_priority = {"CRITICA": "ALTA", "ALTA": "MEDIA", "MEDIA": "BAJA", "BAJA": "BAJA"}[magnitude]
    return reduced_priority, max(10, magnitude_score - 20), "Examen no prioritario; se conserva para revisión completa ordenado por magnitud."

def compare_and_update(reference_csv: str, archivo_rows: list[dict], control_rows: list[dict], sairc_rows: list[dict]):
    if not reference_csv:
        return
    if not (len(archivo_rows) == len(control_rows) == len(sairc_rows)):
        raise ValueError("Las colecciones archivo/control/SAIRC deben tener la misma cantidad de filas.")
    
    try:
        if reference_csv.lower().endswith('.xlsx') or reference_csv.lower().endswith('.xls'):
            df_true = pd.read_excel(reference_csv)
        else:
            try:
                df_true = pd.read_csv(reference_csv, sep=';', encoding='utf-8')
                if len(df_true.columns) <= 1:
                    df_true = pd.read_csv(reference_csv, sep=',', encoding='utf-8')
            except:
                df_true = pd.read_csv(reference_csv, sep=',', encoding='utf-8')
    except Exception as exc:
        print(f"[COMPARADOR] Error al cargar referencia: {exc}")
        return

    cols = list(df_true.columns)
    
    referencia = {}
    for idx, row in df_true.iterrows():
        doc = normalize_document(row.get('Num Doc.'))
        anio = None
        mes = None
        try:
            fecha_m = str(row.get('Fecha de Muestra', '')).strip()
            if fecha_m and fecha_m != 'nan':
                y, m = parse_date(fecha_m)
                if y and m:
                    anio = int(y)
                    mes = int(m)
            
            if not anio or not mes:
                anio = int(row.get('Anio', 0))
                mes = int(row.get('Mes', 0))
        except:
            pass
        
        if doc and anio and mes:
            referencia[(doc, anio, mes)] = row
            
    for a_row, c_row, s_row in zip(archivo_rows, control_rows, sairc_rows):
        doc_val = a_row.get("documento_archivo_pdf")
        fecha_val = a_row.get("fecha_toma_muestra")
        
        if not doc_val:
            val1 = val2 = "Sin documento"
        else:
            doc_str = normalize_document(doc_val)
            
            anio, mes = parse_date(fecha_val)
            if not anio or not mes:
                val1 = val2 = "Fecha invalida o vacia"
            else:
                ref_data = referencia.get((doc_str, anio, mes))
                if ref_data is None:
                    val1 = "No encontrado en CSV"
                    val2 = "No encontrado en CSV"
                else:
                    diffs_g1 = []
                    diffs_g2 = []
                    for sairc_key, ref_key in FIELD_MAP.items():
                        if ref_key not in cols:
                            continue
                        ocr_val = safe_float(s_row.get(sairc_key))
                        ref_val = safe_float(ref_data.get(ref_key))
                        if not numeric_values_equal(ocr_val, ref_val):
                            diff_str = f"{sairc_key} (OCR: {ocr_val}, CSV: {ref_val})"
                            if sairc_key in GROUP1_KEYS:
                                diffs_g1.append(diff_str)
                            else:
                                diffs_g2.append(diff_str)
                    
                    val1 = "; ".join(diffs_g1) if diffs_g1 else "Coincide"
                    val2 = "; ".join(diffs_g2) if diffs_g2 else "Coincide"
        
        a_row["Discrepancias_Urea_Hemo"] = val1
        c_row["Discrepancias_Urea_Hemo"] = val1
        s_row["Discrepancias_Urea_Hemo"] = val1
        a_row["Discrepancias_Resto"] = val2
        c_row["Discrepancias_Resto"] = val2
        s_row["Discrepancias_Resto"] = val2

def _compare_excels_legacy(ocr_excel_path: str, reference_csv: str, output_excel_path: str) -> dict:
    """
    Compara un Excel de OCR (resultados_laboratorio_sairc.xlsx) contra el CSV de SAIRC.
    Mantiene únicamente la hoja SAIRC_Formulario2 en el archivo de salida y añade
    las discrepancias encontradas para todos los valores cuantitativos.
    Retorna un diccionario con estadísticas.
    """
    try:
        if reference_csv.lower().endswith('.xlsx') or reference_csv.lower().endswith('.xls'):
            df_true = pd.read_excel(reference_csv)
        else:
            try:
                df_true = pd.read_csv(reference_csv, sep=';', encoding='utf-8')
                if len(df_true.columns) <= 1:
                    df_true = pd.read_csv(reference_csv, sep=',', encoding='utf-8')
            except:
                df_true = pd.read_csv(reference_csv, sep=',', encoding='utf-8')
    except Exception as exc:
        raise ValueError(f"Error al cargar archivo de referencia: {exc}")

    cols = list(df_true.columns)
    
    referencia = {}
    for idx, row in df_true.iterrows():
        doc = normalize_document(row.get('Num Doc.'))
        
        anio = None
        mes = None
        try:
            # Intentar extraer de Fecha de Muestra si existe
            fecha_m = str(row.get('Fecha de Muestra', '')).strip()
            if fecha_m and fecha_m != 'nan':
                y, m = parse_date(fecha_m)
                if y and m:
                    anio = int(y)
                    mes = int(m)
            
            # Si no, intentar leer de columnas Anio/Mes (legacy)
            if not anio or not mes:
                anio = int(row.get('Anio', 0))
                mes = int(row.get('Mes', 0))
        except: pass
        
        if doc and anio and mes:
            referencia[(doc, anio, mes)] = row

    # Cargar el Excel de OCR
    try:
        from openpyxl import load_workbook
        wb = load_workbook(ocr_excel_path)
    except Exception as exc:
        raise ValueError(f"Error al cargar el Excel de OCR: {exc}")

    # Eliminar hojas innecesarias, quedarnos solo con SAIRC_Formulario2
    sheets = wb.sheetnames
    if 'SAIRC_Formulario2' not in sheets:
        raise ValueError(f"El Excel no contiene la hoja 'SAIRC_Formulario2'. Hojas encontradas: {sheets}")
        
    for sheet_name in sheets:
        if sheet_name != 'SAIRC_Formulario2':
            del wb[sheet_name]
            
    ws = wb['SAIRC_Formulario2']
    headers = [str(c.value).strip() if c.value else '' for c in ws[1]]
    
    col_doc = headers.index('documento_archivo_pdf') if 'documento_archivo_pdf' in headers else -1
    col_fecha = headers.index('fecha_toma_muestra') if 'fecha_toma_muestra' in headers else -1
    
    if col_doc == -1 or col_fecha == -1:
        raise ValueError("No se encontraron las columnas 'documento_archivo_pdf' o 'fecha_toma_muestra' en la hoja.")

    # Mapeo de índices de las columnas cuantitativas
    col_indices = {}
    for ocr_key in FIELD_MAP.keys():
        if ocr_key in headers:
            col_indices[ocr_key] = headers.index(ocr_key)

    new_col1_idx = len(headers) + 1
    new_col2_idx = len(headers) + 2
    ws.cell(row=1, column=new_col1_idx, value='Discrepancias_Urea_Hemo')
    ws.cell(row=1, column=new_col2_idx, value='Discrepancias_Resto')
    
    stats = {'coinciden': 0, 'discrepan': 0, 'no_encontrados': 0, 'errores_fecha': 0, 'sin_documento': 0}
    
    for row_idx in range(2, ws.max_row + 1):
        doc_val = ws.cell(row=row_idx, column=col_doc+1).value
        fecha_val = ws.cell(row=row_idx, column=col_fecha+1).value
        
        if not doc_val:
            ws.cell(row=row_idx, column=new_col1_idx, value="Sin documento")
            ws.cell(row=row_idx, column=new_col2_idx, value="Sin documento")
            stats['sin_documento'] += 1
            continue
            
        doc_str = normalize_document(doc_val)
            
        anio, mes = parse_date(fecha_val)
        if not anio or not mes:
            ws.cell(row=row_idx, column=new_col1_idx, value="Fecha invalida o vacia")
            ws.cell(row=row_idx, column=new_col2_idx, value="Fecha invalida o vacia")
            stats['errores_fecha'] += 1
            continue
            
        ref_data = referencia.get((doc_str, anio, mes))
        if ref_data is None:
            ws.cell(row=row_idx, column=new_col1_idx, value="No encontrado en CSV")
            ws.cell(row=row_idx, column=new_col2_idx, value="No encontrado en CSV")
            stats['no_encontrados'] += 1
            continue
            
        diffs_g1 = []
        diffs_g2 = []
        for ocr_key, ref_key in FIELD_MAP.items():
            if ref_key not in cols or ocr_key not in col_indices:
                continue
            
            ocr_col_idx = col_indices[ocr_key]
            ocr_val = safe_float(ws.cell(row=row_idx, column=ocr_col_idx+1).value)
            ref_val = safe_float(ref_data.get(ref_key))
            
            if not numeric_values_equal(ocr_val, ref_val):
                diff_str = f"{ocr_key} (OCR: {ocr_val}, CSV: {ref_val})"
                if ocr_key in GROUP1_KEYS:
                    diffs_g1.append(diff_str)
                else:
                    diffs_g2.append(diff_str)
                
        if diffs_g1 or diffs_g2:
            stats['discrepan'] += 1
        else:
            stats['coinciden'] += 1

        ws.cell(row=row_idx, column=new_col1_idx, value="; ".join(diffs_g1) if diffs_g1 else "Coincide")
        ws.cell(row=row_idx, column=new_col2_idx, value="; ".join(diffs_g2) if diffs_g2 else "Coincide")
            
    save_workbook_atomic(wb, Path(output_excel_path))
    return stats


def _compare_excels_full(ocr_excel_path: str, reference_csv: str, output_excel_path: str) -> dict:
    """Compara el consolidado con SAIRC y genera un reporte auditable por fila."""
    from copy import copy

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    validation = validate_comparison_inputs(ocr_excel_path, reference_csv)
    reference = _read_reference_table(reference_csv)
    reference_columns = {str(column).strip() for column in reference.columns}
    reference_rows: dict[tuple[str, int, int], object] = {}
    duplicate_keys: set[tuple[str, int, int]] = set()

    for _, row in reference.iterrows():
        document = normalize_document(row.get("Num Doc."))
        year = month = None
        year, month = parse_date(row.get("Fecha de Muestra"))
        if not year or not month:
            try:
                year, month = int(row.get("Anio")), int(row.get("Mes"))
            except (TypeError, ValueError):
                year = month = None
        if document and year and month:
            key = (document, year, month)
            if key in reference_rows:
                duplicate_keys.add(key)
            else:
                reference_rows[key] = row

    workbook = load_workbook(ocr_excel_path)
    worksheet = workbook["SAIRC_Formulario2"]
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in worksheet[1]]
    header_index = {header: index + 1 for index, header in enumerate(headers) if header}
    document_column = header_index.get("documento_archivo_pdf") or header_index.get("documento_identidad")
    sample_date_column = header_index.get("fecha_toma_muestra") or header_index.get("fecha_examen")
    if not document_column or not sample_date_column:
        raise ValueError("No se encontraron columnas de documento y fecha en SAIRC_Formulario2.")

    output_columns = {
        "Discrepancias_Urea_Hemo": len(headers) + 1,
        "Discrepancias_Resto": len(headers) + 2,
        "Estado_validacion": len(headers) + 3,
        "Campos_sin_ocr": len(headers) + 4,
        "Campos_sin_sairc": len(headers) + 5,
    }
    for name, column in output_columns.items():
        existing = header_index.get(name)
        if existing:
            output_columns[name] = existing
        else:
            worksheet.cell(row=1, column=column, value=name)

    stats = {
        "coinciden": 0,
        "discrepan": 0,
        "no_encontrados": 0,
        "errores_fecha": 0,
        "sin_documento": 0,
        "filas_ocr": max(0, worksheet.max_row - 1),
        "filas_referencia": len(reference),
        "filas_con_referencia": 0,
        "referencias_duplicadas": len(duplicate_keys),
        "campos_comparados": 0,
        "campos_sin_ocr": 0,
        "campos_sin_sairc": 0,
        "campos_comparables": 0,
        **validation,
    }

    mapped_columns = [(ocr_key, ref_key, header_index[ocr_key]) for ocr_key, ref_key in FIELD_MAP.items() if ref_key in reference_columns and ocr_key in header_index]
    stats["campos_comparables"] = len(mapped_columns)

    for row_number in range(2, worksheet.max_row + 1):
        document = normalize_document(worksheet.cell(row=row_number, column=document_column).value)
        date_value = worksheet.cell(row=row_number, column=sample_date_column).value
        if not document:
            status = "SIN_DOCUMENTO"
            stats["sin_documento"] += 1
            diffs_group1 = diffs_group2 = []
            missing_ocr = missing_sairc = []
        else:
            year, month = parse_date(date_value)
            if not year or not month:
                status = "FECHA_INVALIDA"
                stats["errores_fecha"] += 1
                diffs_group1 = diffs_group2 = []
                missing_ocr = missing_sairc = []
            else:
                reference_row = reference_rows.get((document, year, month))
                if reference_row is None:
                    status = "NO_ENCONTRADO_SAIRC"
                    stats["no_encontrados"] += 1
                    diffs_group1 = diffs_group2 = []
                    missing_ocr = missing_sairc = []
                else:
                    stats["filas_con_referencia"] += 1
                    diffs_group1, diffs_group2 = [], []
                    missing_ocr, missing_sairc = [], []
                    for ocr_key, ref_key, ocr_column in mapped_columns:
                        ocr_value = safe_float(worksheet.cell(row=row_number, column=ocr_column).value)
                        reference_value = safe_float(reference_row.get(ref_key))
                        if ocr_value is None and reference_value is not None:
                            missing_ocr.append(ocr_key)
                            stats["campos_sin_ocr"] += 1
                        elif ocr_value is not None and reference_value is None:
                            missing_sairc.append(ocr_key)
                            stats["campos_sin_sairc"] += 1
                        elif ocr_value is not None and reference_value is not None:
                            stats["campos_comparados"] += 1
                            if not numeric_values_equal(ocr_value, reference_value):
                                item = f"{ocr_key} (OCR: {ocr_value}, SAIRC: {reference_value})"
                                (diffs_group1 if ocr_key in GROUP1_KEYS else diffs_group2).append(item)
                    if diffs_group1 or diffs_group2:
                        status = "DISCREPANCIA"
                        stats["discrepan"] += 1
                    elif missing_ocr or missing_sairc:
                        status = "FALTAN_DATOS"
                    else:
                        status = "COINCIDE"
                        stats["coinciden"] += 1

        worksheet.cell(row=row_number, column=output_columns["Discrepancias_Urea_Hemo"], value="; ".join(diffs_group1) if diffs_group1 else "Coincide")
        worksheet.cell(row=row_number, column=output_columns["Discrepancias_Resto"], value="; ".join(diffs_group2) if diffs_group2 else "Coincide")
        worksheet.cell(row=row_number, column=output_columns["Estado_validacion"], value=status)
        worksheet.cell(row=row_number, column=output_columns["Campos_sin_ocr"], value=", ".join(missing_ocr))
        worksheet.cell(row=row_number, column=output_columns["Campos_sin_sairc"], value=", ".join(missing_sairc))

    # Mantener el lenguaje visual del Hito 2 y hacer legibles las columnas nuevas.
    source_header = worksheet.cell(row=1, column=1)
    column_widths = {
        "Discrepancias_Urea_Hemo": 58,
        "Discrepancias_Resto": 72,
        "Estado_validacion": 24,
        "Campos_sin_ocr": 38,
        "Campos_sin_sairc": 38,
    }
    for name, column in output_columns.items():
        header_cell = worksheet.cell(row=1, column=column)
        header_cell.fill = copy(source_header.fill) if source_header.fill.fill_type == "solid" else PatternFill("solid", fgColor="1F4E78")
        source_font_color = source_header.font.color
        font_color = "FFFFFF"
        if source_font_color and source_font_color.type == "rgb" and source_font_color.rgb:
            font_color = source_font_color.rgb
        header_cell.font = Font(
            name=source_header.font.name,
            size=source_header.font.size,
            bold=True,
            color=font_color,
        )
        header_cell.border = copy(source_header.border)
        header_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.column_dimensions[get_column_letter(column)].width = column_widths[name]

    status_fills = {
        "COINCIDE": ("E2F0D9", "375623"),
        "DISCREPANCIA": ("F4CCCC", "9C0006"),
        "FALTAN_DATOS": ("FFF2CC", "7F6000"),
        "FECHA_INVALIDA": ("FCE4D6", "9C5700"),
        "NO_ENCONTRADO_SAIRC": ("E7E6E6", "595959"),
        "SIN_DOCUMENTO": ("D9EAD3", "274E13"),
    }
    status_column = output_columns["Estado_validacion"]
    for row_number in range(2, worksheet.max_row + 1):
        status_cell = worksheet.cell(row=row_number, column=status_column)
        fill_color, font_color = status_fills.get(str(status_cell.value), ("FFFFFF", "000000"))
        status_cell.fill = PatternFill("solid", fgColor=fill_color)
        status_cell.font = Font(color=font_color, bold=True)
        status_cell.alignment = Alignment(horizontal="center", vertical="center")

    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.freeze_panes = "A2"

    Path(output_excel_path).parent.mkdir(parents=True, exist_ok=True)
    save_workbook_atomic(workbook, Path(output_excel_path))
    from core.review_state import write_review_sources
    write_review_sources(output_excel_path, ocr_excel_path, reference_csv, [field for field, _ref, _index in mapped_columns])
    return stats


def compare_excels(
    ocr_excel_path: str,
    reference_csv: str,
    output_excel_path: str,
    pdf_root: str | None = None,
    review_pdf_dir: str | None = None,
    review_pdf_operation: str = "copy",
) -> dict:
    """Genera hallazgos priorizados y contexto de valores comparados del Hito 4."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    validation = validate_comparison_inputs(ocr_excel_path, reference_csv)
    reference = _read_reference_table(reference_csv)
    reference_columns = {str(column).strip() for column in reference.columns}
    compare_by_ipress = "Renaes" in reference_columns
    reference_rows: dict[tuple, list] = {}
    for _, row in reference.iterrows():
        document = normalize_document(row.get("Num Doc."))
        year, month = parse_date(row.get("Fecha de Muestra"))
        if not year or not month:
            try:
                year, month = int(row.get("Anio")), int(row.get("Mes"))
            except (TypeError, ValueError):
                year = month = None
        if document and year and month:
            key = (document, year, month, normalize_ipress_code(row.get("Renaes"))) if compare_by_ipress else (document, year, month)
            reference_rows.setdefault(key, []).append(row)
    duplicate_keys = {key for key, rows in reference_rows.items() if len(rows) > 1}

    source_workbook = load_workbook(ocr_excel_path, read_only=True, data_only=True)
    source_sheet = source_workbook["SAIRC_Formulario2"]
    source_rows = source_sheet.iter_rows(values_only=True)
    headers = [str(value).strip() if value is not None else "" for value in next(source_rows)]
    source_rows = list(source_rows)
    header_index = {header: index for index, header in enumerate(headers) if header}
    document_column = header_index.get("documento_archivo_pdf", header_index.get("documento_identidad"))
    sample_date_column = header_index.get("fecha_toma_muestra", header_index.get("fecha_examen"))
    if document_column is None or sample_date_column is None:
        source_workbook.close()
        raise ValueError("No se encontraron columnas de documento y fecha en SAIRC_Formulario2.")

    mapped_columns = [
        (ocr_key, ref_key, header_index[ocr_key])
        for ocr_key, ref_key in FIELD_MAP.items()
        if ref_key in reference_columns and ocr_key in header_index
    ]
    findings: list[dict] = []
    comparison_context: list[dict] = []
    matched_reference_keys: set[tuple] = set()
    stats = {
        **validation,
        "coinciden": 0,
        "discrepan": 0,
        "faltan_datos": 0,
        "no_encontrados": 0,
        "errores_fecha": 0,
        "sin_documento": 0,
        "sin_cui_ipress": 0,
        "filas_ocr": len(source_rows),
        "filas_referencia": len(reference),
        "filas_con_referencia": 0,
        "referencias_duplicadas": len(duplicate_keys),
        "campos_comparados": 0,
        "campos_sin_ocr": 0,
        "campos_sin_sairc": 0,
        "campos_comparables": len(mapped_columns),
        "grupos_paciente_periodo": 0,
        "grupos_con_multiples_pdf": 0,
        "grupos_paneles_separados": 0,
        "comparaciones_fecha_exacta": 0,
        "grupos_referencia_ambigua": 0,
        "campos_referencia_ambigua": 0,
        "correlacion_por_ipress": compare_by_ipress,
        "registros_sairc_sin_resultados": 0,
    }

    def source_value(row: tuple, name: str):
        index = header_index.get(name)
        return row[index] if index is not None and index < len(row) else None

    def append_finding(
        *,
        finding_type: str,
        field: str | None,
        row: tuple | None = None,
        reference_row=None,
        ocr_value: float | None = None,
        sairc_value: float | None = None,
        magnitude: str = "NO_APLICA",
        absolute: float | None = None,
        percentage: float | None = None,
        detail: str = "",
        group_override: str | None = None,
    ) -> None:
        priority, score, reason = _finding_priority(field, magnitude, finding_type)
        document = normalize_document(source_value(row, "documento_archivo_pdf") or source_value(row, "documento_identidad")) if row else ""
        if not document and reference_row is not None:
            document = normalize_document(reference_row.get("Num Doc."))
        date_value = source_value(row, "fecha_toma_muestra") or source_value(row, "fecha_examen") if row else None
        if not date_value and reference_row is not None:
            date_value = reference_row.get("Fecha de Muestra")
        patient = source_value(row, "paciente") if row else None
        if not patient and reference_row is not None:
            patient = reference_row.get("Nombres y apellidos")
        relative_path = source_value(row, "ruta_relativa") if row else None
        source_cui, ipress_name, source_cui_origin = resolve_source_ipress(
            relative_path,
            source_value(row, "archivo") if row else None,
        )
        reference_cui = normalize_ipress_code(reference_row.get("Renaes")) if reference_row is not None else ""
        cui_ipress = source_cui or reference_cui
        laboratory = source_value(row, "laboratorio") if row else None
        establishment_origin = source_cui_origin if ipress_name else "REFERENCIA_SAIRC"
        establishment = ipress_name
        if not establishment and reference_row is not None:
            establishment = reference_row.get("Establecimiento")
        if not establishment:
            establishment_origin = "NO_DISPONIBLE"
        group = "PENALIDAD_HB_UREA" if field in PENALTY_AUDIT_FIELDS else "OTROS"
        if finding_type in {"SAIRC_SIN_PDF", "PDF_SIN_REGISTRO_SAIRC", "REGISTRO_SAIRC_SIN_RESULTADOS"}:
            group = "CORRELACION_REGISTRO_PDF"
        if group_override:
            group = group_override
        findings.append(
            {
                "Prioridad": priority,
                "Puntaje_prioridad": score,
                "Tipo_hallazgo": finding_type,
                "Grupo_control": group,
                "Examen": field or "REGISTRO_COMPLETO",
                "Documento": document,
                "Fecha_periodo": str(date_value or ""),
                "Paciente": patient,
                "CUI_IPRESS": cui_ipress,
                "Establecimiento": establishment,
                "Origen_Establecimiento": establishment_origin,
                "Laboratorio_Emisor": laboratory,
                "Archivo_PDF": source_value(row, "archivo") if row else None,
                "Valor_LLM_OCR": ocr_value,
                "Valor_SAIRC": sairc_value,
                "Diferencia_absoluta": absolute,
                "Diferencia_porcentual": percentage,
                "Magnitud": magnitude,
                "Motivo_prioridad": reason,
                "Interpretacion_control": detail,
                "Ruta_relativa": relative_path,
                "Ruta_PDF_Revision": "",
                "Estado_PDF_Revision": "NO_APLICA",
                "Estado_Hito2": source_value(row, "estado_archivo") if row else None,
            }
        )

    source_groups: dict[tuple, list[tuple]] = {}
    serology_columns = [header_index[name] for name in ("HIV", "VDRL", "HBSAG", "ANTI_HBS", "ANTI_HBC", "HCV", "HTLV") if name in header_index]
    for row in source_rows:
        document = normalize_document(row[document_column])
        if not document:
            stats["sin_documento"] += 1
            append_finding(
                finding_type="SIN_DOCUMENTO",
                field=None,
                row=row,
                detail="El PDF no puede correlacionarse porque no tiene documento utilizable.",
            )
            continue
        year, month = parse_date(row[sample_date_column])
        if not year or not month:
            stats["errores_fecha"] += 1
            append_finding(
                finding_type="FECHA_INVALIDA",
                field=None,
                row=row,
                detail="El PDF no puede correlacionarse porque la fecha de muestra no es válida.",
            )
            continue
        cui_ipress, _ipress_name, _cui_origin = resolve_source_ipress(
            source_value(row, "ruta_relativa"),
            source_value(row, "archivo"),
        )
        if compare_by_ipress and not cui_ipress:
            stats["sin_cui_ipress"] += 1
            append_finding(
                finding_type="CUI_IPRESS_NO_IDENTIFICADO",
                field=None,
                row=row,
                detail="El PDF no puede correlacionarse porque no se identificó su CUI/IPRESS en la ruta ni en el nombre del archivo.",
            )
            continue
        key = (document, year, month, cui_ipress) if compare_by_ipress else (document, year, month)
        source_groups.setdefault(key, []).append(row)

    stats["grupos_paciente_periodo"] = len(source_groups)
    ambiguous_reference_groups: set[tuple] = set()
    for key, group_rows in source_groups.items():
        reference_candidates = reference_rows.get(key)
        if not reference_candidates:
            for row in group_rows:
                stats["no_encontrados"] += 1
                values = "; ".join(
                    f"{field}={row[column]}" for field, _ref, column in mapped_columns if safe_float(row[column]) is not None
                )
                append_finding(
                    finding_type="PDF_SIN_REGISTRO_SAIRC",
                    field=None,
                    row=row,
                    detail=f"PDF con resultados sin registro SAIRC correlacionado. Valores detectados: {values or 'sin valores comparables'}.",
                )
            continue

        matched_reference_keys.add(key)
        stats["filas_con_referencia"] += len(group_rows)
        if len(group_rows) > 1:
            stats["grupos_con_multiples_pdf"] += 1

        def row_exact_date(row: tuple) -> dt.date | None:
            return parse_full_date(row[sample_date_column])

        def reference_exact_date(reference_row) -> dt.date | None:
            return parse_full_date(reference_row.get("Fecha de Muestra"))

        def quantitative_count(row: tuple) -> int:
            return sum(safe_float(row[column]) is not None for _field, _ref, column in mapped_columns)

        representative_row = max(group_rows, key=quantitative_count)
        has_serology_rows = [
            row for row in group_rows if any(column < len(row) and row[column] not in (None, "") for column in serology_columns)
        ]
        has_quantitative_rows = [row for row in group_rows if quantitative_count(row) > 0]
        if has_serology_rows and has_quantitative_rows and any(
            source_value(serology_row, "archivo") != source_value(quantitative_row, "archivo")
            and row_exact_date(serology_row) != row_exact_date(quantitative_row)
            for serology_row in has_serology_rows
            for quantitative_row in has_quantitative_rows
        ):
            stats["grupos_paneles_separados"] += 1

        group_has_difference = False
        group_has_missing = False
        ocr_results_present = {
            ocr_key: next(
                (
                    (safe_float(row[ocr_column]), row)
                    for row in group_rows
                    if safe_float(row[ocr_column]) is not None
                ),
                None,
            )
            for ocr_key, _ref_key, ocr_column in mapped_columns
        }
        ocr_results_present = {field: item for field, item in ocr_results_present.items() if item is not None}
        reference_has_any_result = any(
            safe_float(reference_row.get(ref_key)) is not None
            for reference_row in reference_candidates
            for _ocr_key, ref_key, _ocr_column in mapped_columns
        )
        collapsed_empty_reference = bool(ocr_results_present) and not reference_has_any_result
        if collapsed_empty_reference:
            group_has_missing = True
            stats["registros_sairc_sin_resultados"] += 1
            stats["campos_sin_sairc"] += len(ocr_results_present)
            detected_values = "; ".join(
                f"{field}={value:g}" for field, (value, _row) in ocr_results_present.items()
            )
            append_finding(
                finding_type="REGISTRO_SAIRC_SIN_RESULTADOS",
                field="RESULTADOS_COMPLETOS",
                row=representative_row,
                reference_row=reference_candidates[0],
                group_override=(
                    "PENALIDAD_HB_UREA"
                    if set(ocr_results_present) & PENALTY_AUDIT_FIELDS
                    else "CORRELACION_REGISTRO_PDF"
                ),
                detail=(
                    "Existe una fila SAIRC para el paciente, periodo e IPRESS, pero ninguno de los resultados "
                    f"comparables presentes en el PDF fue registrado. Resultados detectados: {detected_values}."
                ),
            )
        for ocr_key, ref_key, ocr_column in mapped_columns:
            if collapsed_empty_reference:
                break
            ocr_candidates = [
                (safe_float(row[ocr_column]), row, row_exact_date(row))
                for row in group_rows
                if safe_float(row[ocr_column]) is not None
            ]
            reference_values = [
                (safe_float(reference_row.get(ref_key)), reference_row, reference_exact_date(reference_row))
                for reference_row in reference_candidates
                if safe_float(reference_row.get(ref_key)) is not None
            ]

            if not ocr_candidates and reference_values:
                group_has_missing = True
                stats["campos_sin_ocr"] += 1
                sairc_value, reference_row, _reference_date = reference_values[0]
                append_finding(
                    finding_type="SAIRC_CAMPO_SIN_VALOR_PDF",
                    field=ocr_key,
                    row=representative_row,
                    reference_row=reference_row,
                    sairc_value=sairc_value,
                    detail=(
                        f"SAIRC contiene el examen, pero no se encontró en ninguno de los {len(group_rows)} PDF "
                        "consolidados del paciente, periodo e IPRESS."
                    ),
                )
            elif ocr_candidates and not reference_values:
                group_has_missing = True
                stats["campos_sin_sairc"] += 1
                ocr_value, source_row, _source_date = ocr_candidates[0]
                append_finding(
                    finding_type="PDF_CAMPO_SIN_REGISTRO_SAIRC",
                    field=ocr_key,
                    row=source_row,
                    reference_row=reference_candidates[0],
                    ocr_value=ocr_value,
                    detail="Alguno de los PDF consolidados del paciente, periodo e IPRESS contiene el examen, pero SAIRC no tiene valor registrado para ese campo.",
                )
            elif ocr_candidates and reference_values:
                comparison_pairs = []
                unresolved_reference = False
                if len(reference_values) > 1:
                    for reference_item in reference_values:
                        exact_ocr = [
                            ocr_item
                            for ocr_item in ocr_candidates
                            if ocr_item[2] is not None and ocr_item[2] == reference_item[2]
                        ]
                        if exact_ocr:
                            selected = next(
                                (item for item in exact_ocr if numeric_values_equal(item[0], reference_item[0])),
                                exact_ocr[0],
                            )
                            comparison_pairs.append((selected, reference_item))
                            stats["comparaciones_fecha_exacta"] += 1
                            continue
                        equal_ocr = next(
                            (item for item in ocr_candidates if numeric_values_equal(item[0], reference_item[0])),
                            None,
                        )
                        if equal_ocr:
                            comparison_pairs.append((equal_ocr, reference_item))
                        else:
                            unresolved_reference = True
                    if unresolved_reference:
                        ambiguous_reference_groups.add(key)
                        stats["campos_referencia_ambigua"] += 1
                        group_has_missing = True
                else:
                    reference_item = reference_values[0]
                    exact_ocr = [
                        item for item in ocr_candidates if item[2] is not None and item[2] == reference_item[2]
                    ]
                    if exact_ocr:
                        selected_ocr = next(
                            (item for item in exact_ocr if numeric_values_equal(item[0], reference_item[0])),
                            exact_ocr[0],
                        )
                        comparison_pairs.append((selected_ocr, reference_item))
                        stats["comparaciones_fecha_exacta"] += 1
                    else:
                        equal_ocr = next(
                            (item for item in ocr_candidates if numeric_values_equal(item[0], reference_item[0])),
                            None,
                        )
                        if equal_ocr:
                            comparison_pairs.append((equal_ocr, reference_item))
                        elif reference_item[2] is not None:
                            dated_candidates = [item for item in ocr_candidates if item[2] is not None]
                            selected_ocr = min(
                                dated_candidates or ocr_candidates,
                                key=lambda item: abs((item[2] - reference_item[2]).days) if item[2] is not None else 99999,
                            )
                            comparison_pairs.append((selected_ocr, reference_item))
                        else:
                            comparison_pairs.append((ocr_candidates[0], reference_item))

                for selected_ocr, selected_reference in comparison_pairs:
                    ocr_value, source_row, source_date = selected_ocr
                    sairc_value, reference_row, reference_date = selected_reference
                    stats["campos_comparados"] += 1
                    comparison_context.append({
                        "Ruta_relativa": source_value(source_row, "ruta_relativa"),
                        "Documento": normalize_document(source_value(source_row, "documento_archivo_pdf")),
                        "Examen": ocr_key,
                        "Hito_2": ocr_value,
                        "Hito_3": sairc_value,
                        "Estado": "COINCIDE" if numeric_values_equal(ocr_value, sairc_value) else "DIFERENCIA",
                        "Correlacion": "FECHA_EXACTA" if source_date is not None and source_date == reference_date else "PACIENTE_PERIODO_IPRESS",
                    })
                    if numeric_values_equal(ocr_value, sairc_value):
                        continue
                    group_has_difference = True
                    absolute, percentage, magnitude = discrepancy_magnitude(ocr_key, ocr_value, sairc_value)
                    date_note = (
                        f" Comparación por fecha exacta {source_date.isoformat()}."
                        if source_date is not None and source_date == reference_date
                        else " Comparación consolidada por paciente, periodo e IPRESS."
                    )
                    append_finding(
                        finding_type="VALOR_DISCREPANTE",
                        field=ocr_key,
                        row=source_row,
                        reference_row=reference_row,
                        ocr_value=ocr_value,
                        sairc_value=sairc_value,
                        magnitude=magnitude,
                        absolute=absolute,
                        percentage=percentage,
                        detail=f"Los dos orígenes informan valores distintos; revisar PDF y registro SAIRC.{date_note}",
                    )
        if group_has_difference:
            stats["discrepan"] += 1
        elif group_has_missing:
            stats["faltan_datos"] += 1
        else:
            stats["coinciden"] += 1
    stats["grupos_referencia_ambigua"] = len(ambiguous_reference_groups)
    source_workbook.close()

    for key, reference_candidates in reference_rows.items():
        if key in matched_reference_keys:
            continue
        reference_row = reference_candidates[0]
        values = "; ".join(
            f"{ocr_key}={reference_row.get(ref_key)}"
            for ocr_key, ref_key in FIELD_MAP.items()
            if ref_key in reference_columns and safe_float(reference_row.get(ref_key)) is not None
        )
        append_finding(
            finding_type="SAIRC_SIN_PDF",
            field=None,
            reference_row=reference_row,
            detail=f"Registro SAIRC sin informe PDF correlacionado. Valores registrados: {values or 'sin valores comparables'}.",
        )

    for key in duplicate_keys:
        candidates = reference_rows[key]
        reference_row = candidates[0]
        dates = sorted({str(value) for value in (candidate.get("Fecha de Muestra") for candidate in candidates) if value is not None})
        append_finding(
            finding_type="REFERENCIA_DUPLICADA",
            field=None,
            reference_row=reference_row,
            detail=(
                f"La referencia contiene {len(candidates)} registros para el mismo documento, año, mes e IPRESS. "
                f"Fechas: {', '.join(dates) or 'no disponibles'}. Se priorizó fecha exacta; lo no resoluble quedó ambiguo."
            ),
        )

    findings.sort(
        key=lambda item: (
            -int(item["Puntaje_prioridad"]),
            -(float(item["Diferencia_porcentual"]) if item["Diferencia_porcentual"] is not None else -1),
            str(item["Documento"]),
            str(item["Examen"]),
        )
    )

    review_destination = Path(review_pdf_dir) if review_pdf_dir else Path(output_excel_path).parent / "PDFS_PARA_REVISION"
    resolved_pdf_root = Path(pdf_root).resolve() if pdf_root else discover_pdf_root(
        ocr_excel_path,
        (item.get("Ruta_relativa") for item in findings),
    )
    review_stats = {
        "seleccionados": 0,
        "copiados": 0,
        "movidos": 0,
        "ya_preparados": 0,
        "no_encontrados": 0,
        "rutas_invalidas": 0,
        "rutas_recuperadas": 0,
        "rutas_ambiguas": 0,
        "obsoletos_eliminados": 0,
    }
    if resolved_pdf_root:
        review_stats = prepare_review_pdfs(
            findings,
            resolved_pdf_root,
            review_destination,
            operation=review_pdf_operation,
        )
    else:
        for finding in findings:
            if _safe_relative_pdf_path(finding.get("Ruta_relativa")):
                finding["Estado_PDF_Revision"] = "RAIZ_PDF_NO_DETECTADA"
        review_stats["seleccionados"] = len(
            {_safe_relative_pdf_path(item.get("Ruta_relativa")) for item in findings if _safe_relative_pdf_path(item.get("Ruta_relativa"))}
        )

    stats["sairc_sin_pdf"] = sum(item["Tipo_hallazgo"] == "SAIRC_SIN_PDF" for item in findings)
    stats["pdf_sin_sairc"] = sum(item["Tipo_hallazgo"] == "PDF_SIN_REGISTRO_SAIRC" for item in findings)
    stats["registros_sairc_sin_resultados"] = sum(
        item["Tipo_hallazgo"] == "REGISTRO_SAIRC_SIN_RESULTADOS" for item in findings
    )
    stats["hallazgos_total"] = len(findings)
    stats["hallazgos_penalidad"] = sum(item["Grupo_control"] == "PENALIDAD_HB_UREA" for item in findings)
    stats["hallazgos_alta_critica"] = sum(item["Prioridad"] in {"CRITICA", "ALTA"} for item in findings)
    stats["pdf_root"] = str(resolved_pdf_root or "")
    stats["carpeta_pdf_revision"] = str(review_destination)
    stats.update({f"pdf_revision_{key}": value for key, value in review_stats.items()})

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Discrepancias"
    worksheet.sheet_view.showGridLines = False
    worksheet.sheet_view.zoomScale = 85
    headers_out = list(findings[0].keys()) if findings else [
        "Prioridad", "Puntaje_prioridad", "Tipo_hallazgo", "Grupo_control", "Examen", "Documento", "Fecha_periodo",
        "Paciente", "CUI_IPRESS", "Establecimiento", "Origen_Establecimiento", "Laboratorio_Emisor", "Archivo_PDF",
        "Valor_LLM_OCR", "Valor_SAIRC", "Diferencia_absoluta", "Diferencia_porcentual", "Magnitud", "Motivo_prioridad",
        "Interpretacion_control", "Ruta_relativa", "Ruta_PDF_Revision", "Estado_PDF_Revision", "Estado_Hito2",
    ]
    header_columns = {header: get_column_letter(index) for index, header in enumerate(headers_out, start=1)}
    last_column = get_column_letter(len(headers_out))
    detail_header_row = 14
    first_data_row = detail_header_row + 1
    last_data_row = first_data_row + max(len(findings), 1) - 1

    worksheet.merge_cells(f"A1:{last_column}1")
    worksheet["A1"] = "HITO 4 — REPORTE PRIORIZADO DE DISCREPANCIAS Y FALTA DE CORRELACIÓN"
    worksheet["A1"].fill = PatternFill("solid", fgColor="17365D")
    worksheet["A1"].font = Font(color="FFFFFF", bold=True, size=16)
    worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 28
    worksheet.merge_cells(f"A2:{last_column}2")
    worksheet["A2"] = (
        "La prioridad es de auditoría y no constituye rango clínico. La falta de correlación es un indicio para revisión; "
        "no demuestra por sí sola una omisión intencional."
    )
    worksheet["A2"].fill = PatternFill("solid", fgColor="FFF2CC")
    worksheet["A2"].font = Font(color="7F6000", italic=True)
    worksheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.row_dimensions[2].height = 32

    worksheet["A4"], worksheet["B4"] = "Resumen de hallazgos", "Cantidad"
    worksheet["D4"], worksheet["E4"] = "Falta de correlación", "Cantidad"
    worksheet["G4"], worksheet["H4"] = "Consolidación paciente-periodo-IPRESS", "Cantidad"
    summary_left = [
        ("Total de hallazgos", f'=COUNTA(${header_columns["Tipo_hallazgo"]}${first_data_row}:${header_columns["Tipo_hallazgo"]}${last_data_row})'),
        ("Hallazgos en hemoglobina/urea", f'=COUNTIF(${header_columns["Grupo_control"]}${first_data_row}:${header_columns["Grupo_control"]}${last_data_row},"PENALIDAD_HB_UREA")'),
        ("Prioridad crítica", f'=COUNTIF(${header_columns["Prioridad"]}${first_data_row}:${header_columns["Prioridad"]}${last_data_row},"CRITICA")'),
        ("Prioridad alta", f'=COUNTIF(${header_columns["Prioridad"]}${first_data_row}:${header_columns["Prioridad"]}${last_data_row},"ALTA")'),
        ("Diferencias menores", f'=COUNTIFS(${header_columns["Tipo_hallazgo"]}${first_data_row}:${header_columns["Tipo_hallazgo"]}${last_data_row},"VALOR_DISCREPANTE",${header_columns["Magnitud"]}${first_data_row}:${header_columns["Magnitud"]}${last_data_row},"BAJA")'),
        ("Discrepancias de valor", f'=COUNTIF(${header_columns["Tipo_hallazgo"]}${first_data_row}:${header_columns["Tipo_hallazgo"]}${last_data_row},"VALOR_DISCREPANTE")'),
        ("Referencias duplicadas", f'=COUNTIF(${header_columns["Tipo_hallazgo"]}${first_data_row}:${header_columns["Tipo_hallazgo"]}${last_data_row},"REFERENCIA_DUPLICADA")'),
    ]
    summary_right = [
        ("SAIRC sin PDF correlacionado", f'=COUNTIF(${header_columns["Tipo_hallazgo"]}${first_data_row}:${header_columns["Tipo_hallazgo"]}${last_data_row},"SAIRC_SIN_PDF")'),
        ("PDF sin registro SAIRC", f'=COUNTIF(${header_columns["Tipo_hallazgo"]}${first_data_row}:${header_columns["Tipo_hallazgo"]}${last_data_row},"PDF_SIN_REGISTRO_SAIRC")'),
        ("Registro SAIRC sin resultados", f'=COUNTIF(${header_columns["Tipo_hallazgo"]}${first_data_row}:${header_columns["Tipo_hallazgo"]}${last_data_row},"REGISTRO_SAIRC_SIN_RESULTADOS")'),
        ("Campo SAIRC sin valor PDF", f'=COUNTIF(${header_columns["Tipo_hallazgo"]}${first_data_row}:${header_columns["Tipo_hallazgo"]}${last_data_row},"SAIRC_CAMPO_SIN_VALOR_PDF")'),
        ("Campo PDF sin valor SAIRC", f'=COUNTIF(${header_columns["Tipo_hallazgo"]}${first_data_row}:${header_columns["Tipo_hallazgo"]}${last_data_row},"PDF_CAMPO_SIN_REGISTRO_SAIRC")'),
        ("Fecha inválida", f'=COUNTIF(${header_columns["Tipo_hallazgo"]}${first_data_row}:${header_columns["Tipo_hallazgo"]}${last_data_row},"FECHA_INVALIDA")'),
        ("Sin documento", f'=COUNTIF(${header_columns["Tipo_hallazgo"]}${first_data_row}:${header_columns["Tipo_hallazgo"]}${last_data_row},"SIN_DOCUMENTO")'),
        ("Sin CUI/IPRESS", f'=COUNTIF(${header_columns["Tipo_hallazgo"]}${first_data_row}:${header_columns["Tipo_hallazgo"]}${last_data_row},"CUI_IPRESS_NO_IDENTIFICADO")'),
        ("Alta/crítica penalizable", f'=COUNTIFS(${header_columns["Grupo_control"]}${first_data_row}:${header_columns["Grupo_control"]}${last_data_row},"PENALIDAD_HB_UREA",${header_columns["Prioridad"]}${first_data_row}:${header_columns["Prioridad"]}${last_data_row},"CRITICA")+COUNTIFS(${header_columns["Grupo_control"]}${first_data_row}:${header_columns["Grupo_control"]}${last_data_row},"PENALIDAD_HB_UREA",${header_columns["Prioridad"]}${first_data_row}:${header_columns["Prioridad"]}${last_data_row},"ALTA")'),
    ]
    for offset, (label, formula) in enumerate(summary_left, start=5):
        worksheet.cell(offset, 1, label)
        worksheet.cell(offset, 2, formula)
    for offset, (label, formula) in enumerate(summary_right, start=5):
        worksheet.cell(offset, 4, label)
        worksheet.cell(offset, 5, formula)
    consolidation_summary = [
        ("Grupos paciente-periodo-IPRESS", stats["grupos_paciente_periodo"]),
        ("Grupos con múltiples PDF", stats["grupos_con_multiples_pdf"]),
        ("Paneles separados", stats["grupos_paneles_separados"]),
        ("Comparaciones por fecha exacta", stats["comparaciones_fecha_exacta"]),
        ("Referencias ambiguas", stats["grupos_referencia_ambigua"]),
        ("Campos ambiguos", stats["campos_referencia_ambigua"]),
    ]
    for offset, (label, value) in enumerate(consolidation_summary, start=5):
        worksheet.cell(offset, 7, label)
        worksheet.cell(offset, 8, value)
    for cell in [worksheet["A4"], worksheet["B4"], worksheet["D4"], worksheet["E4"], worksheet["G4"], worksheet["H4"]]:
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center")
    for row_number in range(5, 14):
        worksheet.cell(row_number, 2).number_format = "#,##0"
        worksheet.cell(row_number, 5).number_format = "#,##0"
        worksheet.cell(row_number, 8).number_format = "#,##0"

    for column, header in enumerate(headers_out, start=1):
        cell = worksheet.cell(detail_header_row, column, header)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[detail_header_row].height = 34
    pdf_urls = {}
    for row_number, finding in enumerate(findings, start=first_data_row):
        for column, header in enumerate(headers_out, start=1):
            cell = worksheet.cell(row_number, column, finding.get(header))
            if header == "Ruta_PDF_Revision" and finding.get(header):
                cell.value = "Abrir PDF"
                pdf_path = str(finding[header])
                if pdf_path not in pdf_urls:
                    pdf_urls[pdf_path] = Path(pdf_path).resolve().as_uri()
                cell.hyperlink = pdf_urls[pdf_path]
                cell.font = Font(color="0563C1", underline="single")
                cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [12, 10, 34, 24, 22, 14, 16, 30, 12, 38, 20, 28, 32, 15, 15, 16, 18, 12, 48, 70, 55, 16, 22, 22]
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width
    for row_number in range(first_data_row, last_data_row + 1):
        priority_cell = worksheet.cell(row_number, 1)
        colors = {
            "CRITICA": ("F4CCCC", "9C0006"),
            "ALTA": ("FCE4D6", "9C5700"),
            "MEDIA": ("FFF2CC", "7F6000"),
            "BAJA": ("E2F0D9", "375623"),
        }
        fill_color, font_color = colors.get(str(priority_cell.value), ("FFFFFF", "000000"))
        priority_cell.fill = PatternFill("solid", fgColor=fill_color)
        priority_cell.font = Font(color=font_color, bold=True)
        priority_cell.alignment = Alignment(horizontal="center")
        worksheet.cell(row_number, 2).number_format = "0"
        worksheet.cell(row_number, 13).number_format = "0.0000"
        worksheet.cell(row_number, 14).number_format = "0.00"

    worksheet.freeze_panes = f"A{first_data_row}"
    worksheet.auto_filter.ref = f"A{detail_header_row}:{last_column}{last_data_row}"
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    context_sheet = workbook.create_sheet("Valores_comparados")
    context_columns = ["Ruta_relativa", "Documento", "Examen", "Hito_2", "Hito_3", "Estado", "Correlacion"]
    context_sheet.append(context_columns)
    review_paths = {item.get("Ruta_relativa") for item in findings}
    for item in comparison_context:
        if item.get("Ruta_relativa") in review_paths:
            context_sheet.append([item.get(column) for column in context_columns])
    context_sheet.freeze_panes = "A2"
    context_sheet.auto_filter.ref = context_sheet.dimensions
    Path(output_excel_path).parent.mkdir(parents=True, exist_ok=True)
    save_workbook_atomic(workbook, Path(output_excel_path))
    from core.review_state import digest, write_review_sources
    pdf_sources = {}
    if resolved_pdf_root and review_pdf_operation == "copy":
        destination_absolute = review_destination.resolve()
        seen_paths = set()
        manifest_path = review_destination / "copias_revision.json"
        owned = json.loads(manifest_path.read_text(encoding="utf-8")).get("copies", {}) if manifest_path.is_file() else {}
        for finding in findings:
            path = str(finding.get("Ruta_PDF_Revision") or "")
            if not path or path in seen_paths:
                continue
            seen_paths.add(path)
            target = Path(path).resolve()
            if target.is_relative_to(destination_absolute) and target.is_file():
                relative = target.relative_to(destination_absolute)
                original = (resolved_pdf_root / relative).resolve()
                if original.is_relative_to(resolved_pdf_root):
                    pdf_sources[str(target)] = {"path": str(original), "sha256": owned.get(relative.as_posix()) or digest(target)}
    write_review_sources(output_excel_path, ocr_excel_path, reference_csv, [field for field, _ref, _index in mapped_columns], pdf_sources)
    return stats
