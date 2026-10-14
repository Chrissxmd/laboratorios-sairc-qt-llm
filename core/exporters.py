from __future__ import annotations

import csv
import logging
import os
import tempfile
from datetime import datetime
from pathlib import Path

log = logging.getLogger("lab_extractor.exporters")


def write_csv_dicts(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    headers: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in headers and not k.startswith("_"):
                headers.append(k)
    if not headers:
        headers = ["sin_datos"]
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for r in rows:
            writer.writerow({k: r.get(k) for k in headers})


def write_csv(path: Path, rows: list[dict]) -> None:
    headers: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in headers:
                headers.append(k)
    if not headers:
        headers = ["sin_datos"]
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def split_sairc_rows(sairc_rows: list[dict]) -> tuple[list[dict], list[dict], list[dict]]:
    cargable: list[dict] = []
    revision: list[dict] = []
    no_cargable: list[dict] = []
    for r in sairc_rows:
        estado = str(r.get("estado_archivo") or "").upper()
        puede = str(r.get("puede_cargarse_sairc") or "").upper()
        if puede == "SI" and estado in {"OK_CARGABLE", "OK_CON_ALERTA_CLINICA", "OK", "OK_CON_REGLA"}:
            cargable.append(r)
        elif estado == "SAIRC_REVISION":
            revision.append(r)
        else:
            no_cargable.append(r)
    return cargable, revision, no_cargable


def write_sairc_csvs(output_dir: Path, sairc_rows: list[dict]) -> None:
    cargable, revision, no_cargable = split_sairc_rows(sairc_rows)
    write_csv(output_dir / "sairc_formulario2.csv", sairc_rows)
    write_csv(output_dir / "sairc_cargable.csv", cargable)
    write_csv(output_dir / "sairc_revision.csv", revision)
    write_csv(output_dir / "sairc_no_cargable.csv", no_cargable)
    write_csv(output_dir / "sairc_todos_con_estado.csv", sairc_rows)


def build_resumen_ipress(archivo_rows: list[dict]) -> list[dict]:
    acc: dict[str, dict] = {}
    for r in archivo_rows:
        key = r.get("carpeta_origen") or "(raiz)"
        row = acc.setdefault(
            key,
            {
                "carpeta_origen": key,
                "total_archivos": 0,
                "ok_cargable": 0,
                "revision": 0,
                "no_cargable_identidad": 0,
                "sin_doc_pdf": 0,
                "errores_tecnicos": 0,
                "recomendacion_llm_vision": "NO",
                "motivo_recomendacion_llm_vision": None,
                "costo_estimado_usd": 0.0,
                "segundos_total": 0.0,
            },
        )
        row["total_archivos"] += 1
        estado = str(r.get("estado_archivo") or "").upper()
        est_id = str(r.get("estado_identidad") or "").upper()
        if estado in {"OK_CARGABLE", "OK_CON_ALERTA_CLINICA", "OK", "OK_CON_REGLA"}:
            row["ok_cargable"] += 1
        elif estado == "NO_CARGABLE_IDENTIDAD" or est_id == "NO_CARGABLE_IDENTIDAD":
            row["no_cargable_identidad"] += 1
        elif "SIN_DOCUMENTO" in est_id or str(r.get("coincidencia_documento") or "").upper() == "SIN_DOC_OCR":
            row["sin_doc_pdf"] += 1
        elif estado or r.get("status") == "success":
            row["revision"] += 1
        else:
            row["errores_tecnicos"] += 1
        if str(r.get("recomendacion_llm_vision") or "").upper() == "SI":
            row["recomendacion_llm_vision"] = "SI"
            row["motivo_recomendacion_llm_vision"] = r.get("motivo_recomendacion_llm_vision")
        try:
            row["costo_estimado_usd"] += float(r.get("costo_estimado_usd") or 0)
        except Exception:
            pass
        try:
            row["segundos_total"] += float(r.get("seconds") or 0)
        except Exception:
            pass
    out = []
    for row in acc.values():
        total = row["total_archivos"] or 1
        row["costo_estimado_usd"] = round(row["costo_estimado_usd"], 6)
        row["segundos_promedio"] = round(row["segundos_total"] / total, 2)
        row["segundos_total"] = round(row["segundos_total"], 2)
        out.append(row)
    return sorted(out, key=lambda x: str(x.get("carpeta_origen")))


def build_llm_vision_recommendations(archivo_rows: list[dict], pendientes: list[dict]) -> list[dict]:
    acc: dict[str, dict] = {}
    for r in archivo_rows:
        key = r.get("carpeta_origen") or "(raiz)"
        row = acc.setdefault(
            key,
            {
                "carpeta_origen": key,
                "total_archivos": 0,
                "archivos_con_inconsistencias": 0,
                "pendientes_alta": 0,
                "errores_tecnicos": 0,
                "identidad_no_cargable": 0,
                "sin_documento_pdf": 0,
                "recomendar_llm_vision": "NO",
                "accion_recomendada": None,
                "motivos": "",
            },
        )
        row["total_archivos"] += 1
        estado = str(r.get("estado_archivo") or "").upper()
        est_id = str(r.get("estado_identidad") or "").upper()
        status = str(r.get("status") or "").upper()
        motivos = set(filter(None, str(row.get("motivos") or "").split("; ")))
        if status and status != "SUCCESS":
            row["errores_tecnicos"] += 1
            motivos.add("errores_tecnicos")
        if estado not in {"OK_CARGABLE", "OK_CON_ALERTA_CLINICA", "OK", "OK_CON_REGLA"}:
            row["archivos_con_inconsistencias"] += 1
            motivos.add(estado.lower() or "estado_no_cargable")
        if estado == "NO_CARGABLE_IDENTIDAD" or est_id == "NO_CARGABLE_IDENTIDAD":
            row["identidad_no_cargable"] += 1
            motivos.add("identidad_no_cargable")
        if "SIN_DOCUMENTO" in est_id or str(r.get("coincidencia_documento") or "").upper() == "SIN_DOC_OCR":
            row["sin_documento_pdf"] += 1
            motivos.add("documento_no_detectado_en_pdf")
        row["motivos"] = "; ".join(sorted(motivos))

    for p in pendientes:
        key = p.get("carpeta_origen") or "(raiz)"
        row = acc.setdefault(
            key,
            {
                "carpeta_origen": key,
                "total_archivos": 0,
                "archivos_con_inconsistencias": 0,
                "pendientes_alta": 0,
                "errores_tecnicos": 0,
                "identidad_no_cargable": 0,
                "sin_documento_pdf": 0,
                "recomendar_llm_vision": "NO",
                "accion_recomendada": None,
                "motivos": "",
            },
        )
        motivos = set(filter(None, str(row.get("motivos") or "").split("; ")))
        if str(p.get("severidad") or "").upper() == "ALTA":
            row["pendientes_alta"] += 1
            motivos.add("pendientes_alta")
        motivo = str(p.get("motivo") or "")
        if motivo:
            motivos.add(motivo)
        row["motivos"] = "; ".join(sorted(motivos))

    for row in acc.values():
        has_inconsistencies = any(
            int(row.get(k) or 0) > 0
            for k in ["archivos_con_inconsistencias", "pendientes_alta", "errores_tecnicos", "identidad_no_cargable", "sin_documento_pdf"]
        )
        if has_inconsistencies:
            row["recomendar_llm_vision"] = "SI"
            row["accion_recomendada"] = "Reprocesar todos los PDFs de este laboratorio/carpeta con lector LLM vision."
    return sorted(acc.values(), key=lambda x: str(x.get("carpeta_origen")))


def _write_sheet(wb, title: str, rows: list[dict]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    ws = wb.active if wb.active.title == "Sheet" else wb.create_sheet(title)
    ws.title = title[:31]
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 85
    if not rows:
        ws.append(["sin_datos"])
        ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
        ws["A1"].font = Font(color="FFFFFF", bold=True)
        return
    headers: list[str] = []
    for r in rows:
        for k in r.keys():
            if k not in headers:
                headers.append(k)
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    for r in rows:
        ws.append([r.get(h) for h in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        try:
            max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 45)
        except Exception:
            pass


def save_workbook_atomic(workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as stream:
        temporary = Path(stream.name)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def write_excel(
    output_dir: Path,
    resumen_rows: list[dict],
    control_rows: list[dict],
    normalizados: list[dict],
    sairc_rows: list[dict],
    archivo_rows: list[dict],
    pendientes: list[dict],
    raw_rows: list[dict],
    costos_rows: list[dict],
    error_rows: list[dict],
) -> Path:
    try:
        from openpyxl import Workbook

        wb = Workbook()
        sairc_cargable, sairc_revision, sairc_no_cargable = split_sairc_rows(sairc_rows)
        control_identidad = [
            {
                k: r.get(k)
                for k in [
                    "ruta_relativa",
                    "archivo",
                    "carpeta_origen",
                    "paciente",
                    "documento_archivo_pdf",
                    "documento_archivo_normalizado",
                    "documento_identidad",
                    "documento_pdf_normalizado",
                    "coincidencia_documento",
                    "estado_identidad",
                    "requiere_verificacion_identidad",
                    "requiere_verificacion_resultados",
                    "requiere_verificacion_humana",
                    "motivo_identidad",
                    "firma_ctmp",
                    "sugerencia_llm",
                    "documentos_detectados",
                    "puede_cargarse_sairc",
                ]
            }
            for r in control_rows
        ]
        alertas_calidad = []
        for r in pendientes:
            alertas_calidad.append(r)
        for r in normalizados:
            if (
                str(r.get("estado_control") or "").upper() == "OK_CON_REGLA"
                or str(r.get("estado_produccion") or "").upper() == "OK_CON_ALERTA_CLINICA"
            ):
                alertas_calidad.append(r)
        resumen_ipress = build_resumen_ipress(archivo_rows)
        recomendaciones_llm = build_llm_vision_recommendations(archivo_rows, pendientes)
        _write_sheet(wb, "Resumen", resumen_rows)
        _write_sheet(wb, "Resumen_IPRESS", resumen_ipress)
        _write_sheet(wb, "Recomendacion_LLM_Vision", recomendaciones_llm)
        _write_sheet(wb, "Control_Identidad", control_identidad)
        _write_sheet(wb, "Control_Calidad", control_rows)
        _write_sheet(wb, "SAIRC_Cargable", sairc_cargable)
        _write_sheet(wb, "SAIRC_Revision", sairc_revision)
        _write_sheet(wb, "SAIRC_No_Cargable", sairc_no_cargable)
        # Hoja base del Hito 4: conserva el universo completo. La selección segura
        # para carga se mantiene explícitamente en SAIRC_Cargable.
        _write_sheet(wb, "SAIRC_Formulario2", sairc_rows)
        _write_sheet(wb, "SAIRC_Todos_Con_Estado", sairc_rows)
        _write_sheet(wb, "Resultados_Normalizados", normalizados)
        _write_sheet(wb, "Archivos", archivo_rows)
        pendientes_alta = [
            r for r in pendientes if str(r.get("severidad") or "").upper() == "ALTA" or "identidad" in str(r.get("motivo") or "").lower()
        ]
        _write_sheet(wb, "Pendientes_Alta", pendientes_alta)
        _write_sheet(wb, "Pendientes_Revision", pendientes)
        _write_sheet(wb, "Alertas_Calidad", alertas_calidad)
        _write_sheet(wb, "Raw_Gemini", raw_rows)
        _write_sheet(wb, "Costos_Latencia", costos_rows)
        _write_sheet(wb, "Errores", error_rows)
        excel_path = output_dir / "resultados_laboratorio_sairc.xlsx"
        try:
            save_workbook_atomic(wb, excel_path)
        except PermissionError:
            fallback = output_dir / f"resultados_laboratorio_sairc_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
            save_workbook_atomic(wb, fallback)
            log.warning("Excel principal bloqueado. Se guardo copia alternativa: %s", fallback)
            return fallback.resolve()
        return excel_path.resolve()
    except Exception as exc:
        raise RuntimeError(f"Exportación incompleta: no se pudo generar el Excel en {output_dir}: {exc}") from exc
