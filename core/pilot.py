"""Prueba secuencial de tres modelos sobre un conjunto invariable de PDF."""
import hashlib
import json
import time
from dataclasses import replace
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from core.ai import MISTRAL_PILOT_MODELS, _get, normalize_cached_lab_json, validate_reader_config, validate_remote_model
from core.exporters import save_workbook_atomic
from core.persistence import write_json_atomic
from core.pdf_utils import PDF_IMAGE_VERSION
from core.processor import _resolve_ai_reprocess_source, _resolve_state_dir, process_folder, reprocess_ai_cache
from core.sairc_normalizer import normalize_ai_output
from core.workspace import get_hito_paths


def _dataset(source):
    files = sorted(source.rglob("*.pdf"))
    result = {}
    for path in files:
        with path.open("rb") as stream:
            result[path.relative_to(source).as_posix()] = hashlib.file_digest(stream, "sha256").hexdigest()
    return result


def _report(root, summaries, readings):
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    columns = ["modelo", "pdf_total", "exitos", "errores", "reutilizados", "tokens_entrada", "tokens_salida", "costo_acumulado_usd", "segundos_esta_ejecucion", "excel_sairc"]
    summary.append(columns)
    for row in summaries:
        summary.append([row.get(column) for column in columns])
    details = workbook.create_sheet("Lecturas_comparadas")
    details.append(["PDF", "Campo", *MISTRAL_PILOT_MODELS, "Coinciden_los_tres"])
    for pdf in sorted(readings):
        model_values = readings[pdf]
        fields = sorted({field for values in model_values.values() for field in values})
        for field in fields:
            values = [model_values.get(model, {}).get(field) for model in MISTRAL_PILOT_MODELS]
            agree = "SIN_DATO" if all(value is None for value in values) else ("SI" if len(model_values) == 3 and values[0] == values[1] == values[2] else "NO")
            # Todos los campos se muestran como valores, nunca como fórmulas del proveedor.
            details.append([pdf, field, *values, agree])
    note = workbook.create_sheet("Interpretacion")
    note.append(["Coincidencia entre modelos no demuestra exactitud. Revisar valores contra el PDF original."])
    note.append(["Costo acumulado incluye respuestas reutilizadas; el tiempo corresponde a esta ejecución. Tarifas estándar referenciales, sin descuento por caché del proveedor."])
    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for cells in sheet:
            for cell in cells:
                if cell.data_type == "f":
                    cell.data_type = "s"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(45, max(20, len(str(column[0].value or "")) + 2))
    target = root / "comparacion_modelos.xlsx"
    save_workbook_atomic(workbook, target)
    return target


def run_mistral_pilot(source_folder, workspace, config, *, authorize_three_models=False, threads=2, progress_cb=None, pause_event=None, cancel_event=None):
    validate_reader_config(config)
    if config.provider != "mistral" or config.reader_mode != "vision" or not authorize_three_models:
        raise ValueError("Selecciona Mistral/Visión y autoriza expresamente la prueba con los tres modelos.")
    source = Path(source_folder).expanduser().resolve()
    root = Path(workspace).expanduser().resolve() / "PRUEBA_MISTRAL_3_MODELOS"
    if not source.is_dir():
        raise ValueError("Selecciona la carpeta de PDF de la prueba anterior.")
    if root.is_relative_to(source) or source.is_relative_to(root):
        raise ValueError("La salida de la prueba debe estar separada de la carpeta de PDF.")
    dataset = _dataset(source)
    if not dataset:
        raise ValueError("No hay PDF en la carpeta seleccionada.")
    manifest = root / "documentos_prueba.json"
    if manifest.exists() and json.loads(manifest.read_text(encoding="utf-8")).get("pdf_image_version") != PDF_IMAGE_VERSION:
        raise ValueError("Esta prueba fue creada con imágenes anteriores. Selecciona una nueva raíz de resultados para conservarla y probar con fondo blanco.")
    if manifest.exists() and json.loads(manifest.read_text(encoding="utf-8"))["sha256"] != dataset:
        raise ValueError("El conjunto de PDF cambió. Elige otra raíz de resultados para una prueba nueva.")
    if cancel_event and cancel_event.is_set():
        return {"cancelled": True, "models": [], "report": ""}
    # Comprobar los tres modelos antes de la primera lectura facturable.
    catalog = _get("https://api.mistral.ai/v1/models", headers={"Authorization": f"Bearer {config.api_key}"}, timeout=min(config.timeout_seconds, 30))
    for model in MISTRAL_PILOT_MODELS:
        validate_remote_model(replace(config, model=model), catalog=catalog)
    root.mkdir(parents=True, exist_ok=True)
    write_json_atomic(manifest, {"input_dir": str(source), "sha256": dataset, "models": list(MISTRAL_PILOT_MODELS), "pdf_image_version": PDF_IMAGE_VERSION})
    summaries, readings = [], {pdf: {} for pdf in dataset}
    report = ""
    for index, model in enumerate(MISTRAL_PILOT_MODELS):
        if cancel_event and cancel_event.is_set():
            break
        if _dataset(source) != dataset:
            raise ValueError("Los PDF cambiaron durante la prueba; se detuvo para conservar comparabilidad.")
        paths = get_hito_paths(root / model)
        started = time.monotonic()
        def progress(payload, stage_index=index, name=model):
            if progress_cb:
                progress_cb({**payload, "percent": int((stage_index * 100 + payload.get("percent", 0)) / 3), "message": f"{name}: {payload.get('message', '')}"})
        extracted = process_folder(str(source), replace(config, model=model), threads=threads, output_dir=str(paths.hito1), progress_cb=progress, pause_event=pause_event, cancel_event=cancel_event)
        if _dataset(source) != dataset:
            raise ValueError("Los PDF cambiaron durante la lectura; revisa los resultados y utiliza una prueba nueva.")
        excel = ""
        if extracted.success:
            consolidated = reprocess_ai_cache(str(paths.hito1), str(paths.hito2), state_dir=str(_resolve_state_dir(paths.hito1)))
            excel = consolidated.excel_path
        _, raw_dir = _resolve_ai_reprocess_source(paths.hito1)
        for raw_path in raw_dir.glob("*.json"):
            state_path = _resolve_state_dir(paths.hito1) / f"{raw_path.stem}.meta.json"
            if not state_path.is_file() or json.loads(state_path.read_text(encoding="utf-8")).get("status") != "success":
                continue
            raw = normalize_cached_lab_json(json.loads(raw_path.read_text(encoding="utf-8")))
            relative = raw.get("_source", {}).get("ruta_relativa")
            if relative in readings:
                normalized = normalize_ai_output(raw, ruta_relativa=relative, carpeta_origen=str(Path(relative).parent))
                fields = {"Estado": "EXITO", "Documento_PDF": raw.get("documento_identidad"), "Fecha_toma_PDF": raw.get("fecha_toma_muestra"), "Fecha_examen_PDF": raw.get("fecha_examen"), "Laboratorio_PDF": raw.get("laboratorio"), "Respuesta_JSON": str(raw_path.resolve()), **normalized.get("sairc", {})}
                original_results = {}
                for item in raw.get("resultados", []):
                    original_results.setdefault(str(item.get("analisis", "")), []).append(item)
                fields.update({f"Lectura original: {analysis}": json.dumps(items, ensure_ascii=False, sort_keys=True) for analysis, items in original_results.items()})
                readings[relative][model] = fields
        for relative in dataset:
            readings[relative].setdefault(model, {"Estado": "ERROR_O_NO_PROCESADO"})
        summaries.append({"modelo": model, "pdf_total": extracted.total, "exitos": extracted.success, "errores": extracted.errors, "reutilizados": extracted.skipped_success, "tokens_entrada": extracted.total_prompt_tokens, "tokens_salida": extracted.total_output_tokens, "costo_acumulado_usd": extracted.total_cost_estimated_usd, "segundos_esta_ejecucion": round(time.monotonic() - started, 2), "excel_sairc": excel})
        write_json_atomic(root / "resumen_modelos.json", summaries)
        report = str(_report(root, summaries, readings))
    return {"cancelled": bool(cancel_event and cancel_event.is_set()), "models": summaries, "report": report, "output_dir": str(root)}
