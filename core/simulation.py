"""Hito 3 sintético compatible y comparación real Hito 2–Hito 3, sin IA/SQL."""
from datetime import datetime
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from core.comparador import FIELD_MAP, compare_excels, normalize_document, parse_full_date, resolve_source_ipress, safe_float
from core.exporters import save_workbook_atomic
from core.persistence import write_json_atomic
from core.processor import reprocess_ai_cache
from core.review_state import write_review_sources
from core.workspace import get_hito_paths


def simulate_review(cache_dir: str, pdf_root: str, output_root: str, progress_cb=None) -> dict:
    root = Path(pdf_root).resolve()
    if not pdf_root or not root.is_dir():
        raise ValueError("Selecciona la carpeta de PDF originales para la simulación.")
    pdfs = sorted(path for path in root.rglob("*") if path.is_file() and path.suffix.lower() == ".pdf" and path.resolve().is_relative_to(root))
    if not pdfs:
        raise ValueError("La carpeta seleccionada no contiene PDF.")
    output = Path(output_root).resolve() / "SIMULACION_VISOR" / datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    if output.is_relative_to(Path(cache_dir).resolve()) or output.is_relative_to(root):
        raise ValueError("Los resultados de simulación deben quedar fuera de las lecturas y PDF originales.")
    paths = get_hito_paths(output)
    if progress_cb:
        progress_cb({"message": "Reconstruyendo Hito 2 y creando Hito 3 sintético compatible…"})
    summary = reprocess_ai_cache(cache_dir, str(paths.hito2))
    if not summary.success:
        raise ValueError("No hay respuestas IA disponibles para generar el Hito 2.")
    book = load_workbook(summary.excel_path, read_only=True, data_only=True)
    try:
        values = book["SAIRC_Formulario2"].iter_rows(values_only=True)
        headers = next(values)
        source_rows = [dict(zip(headers, row)) for row in values]
    finally:
        book.close()
    records, groups = {}, {}
    by_ipress = any(resolve_source_ipress(row.get("ruta_relativa"), row.get("archivo"))[0] for row in source_rows)
    for row in source_rows:
        relative = str(row.get("ruta_relativa") or "").replace("\\", "/").casefold()
        records.setdefault(relative, []).append(row)
        document = normalize_document(row.get("documento_archivo_pdf") or row.get("documento_identidad"))
        date = parse_full_date(row.get("fecha_toma_muestra") or row.get("fecha_examen"))
        cui = resolve_source_ipress(row.get("ruta_relativa"), row.get("archivo"))[0]
        if not document or not date:
            continue
        key = (document, date.isoformat(), cui if by_ipress else "")
        reference_row = groups.setdefault(key, {"Num Doc.": document, "Fecha de Muestra": date.isoformat(), **({"Renaes": cui} if by_ipress else {})})
        for field, target in FIELD_MAP.items():
            value = safe_float(row.get(field))
            if value is not None:
                reference_row.setdefault(target, value)
    reference = Workbook()
    sheet = reference.active
    sheet.title = "Referencia_SINTETICA"
    columns = ["Num Doc.", "Fecha de Muestra", *(["Renaes"] if by_ipress else []), *FIELD_MAP.values()]
    sheet.append(columns)
    scenarios = {"coincidencia_control": 0, "decimal_desplazado": 0, "dato_omitido": 0, "valor_alterado": 0, "referencia_duplicada": 0}
    for index, row in enumerate(groups.values()):
        available = [target for target in FIELD_MAP.values() if row.get(target) is not None]
        if available:
            target = available[0]
            pattern = index % 4
            # El primer caso también produce una diferencia para pruebas de visor pequeñas.
            if pattern == 0:
                row[target] = row[target] + 1
                scenarios["valor_alterado"] += 1
            elif pattern == 1:
                row[target] = row[target] * 10 if row[target] else 1
                scenarios["decimal_desplazado"] += 1
            elif pattern == 2:
                row[target] = None
                scenarios["dato_omitido"] += 1
            else:
                scenarios["coincidencia_control"] += 1
        sheet.append([row.get(column) for column in columns])
        if index > 0 and index % 13 == 0:
            sheet.append([row.get(column) for column in columns])
            scenarios["referencia_duplicada"] += 1
    warning = reference.create_sheet("SIMULACION_NO_CARGAR")
    warning.append(["Referencia ficticia derivada de Hito 2 con cambios controlados. No válida para auditoría ni carga SAIRC."])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    reference_path = paths.hito3 / "Referencia_SIMULADA_NO_CARGAR.xlsx"
    save_workbook_atomic(reference, reference_path)
    report_path = paths.hito4 / "reporte_SIMULADO_visor.xlsx"
    if progress_cb:
        progress_cb({"message": "Ejecutando el comparador real Hito 2 ↔ Hito 3…"})
    stats = compare_excels(summary.excel_path, str(reference_path), str(report_path), pdf_root=str(root), review_pdf_operation="copy")
    report = load_workbook(report_path)
    sheet = report["Discrepancias"]
    output_headers = [cell.value for cell in sheet[14]]
    header_index = {name: index for index, name in enumerate(output_headers)}
    relative_column = header_index["Ruta_relativa"]
    represented = {str(row[relative_column] or "").replace("\\", "/").casefold() for row in sheet.iter_rows(min_row=15, values_only=True)}
    missing, controls = 0, 0
    for pdf in pdfs:
        relative = pdf.relative_to(root).as_posix()
        candidates = records.get(relative.casefold(), [])
        if not candidates:
            missing += 1
        if relative.casefold() in represented:
            continue
        control = {"Prioridad": "BAJA", "Tipo_hallazgo": "SIMULACION_COINCIDE" if candidates else "SIMULACION_SIN_LECTURA", "Examen": "CONTROL_VISOR", "Documento": (candidates[0].get("documento_archivo_pdf") if candidates else ""), "Ruta_relativa": relative, "Ruta_PDF_Revision": str(pdf.resolve()), "Interpretacion_control": "Control de navegación sin discrepancia detectada en los campos comparados; no es validación clínica." if candidates else "Sin lectura del Hito 1; incluido para comprobar el visor. Derivar para resolver."}
        sheet.append([control.get(name) for name in output_headers])
        sheet.cell(sheet.max_row, header_index["Ruta_PDF_Revision"] + 1).hyperlink = pdf.resolve().as_uri()
        controls += 1
    sheet["A1"] = "SIMULACIÓN — COMPARACIÓN REAL HITO 2 ↔ HITO 3 CON REFERENCIA FICTICIA"
    sheet["A2"] = "Datos sintéticos de referencia. Diferencias introducidas para probar la revisión. No válidos para auditoría ni carga SAIRC. Los controles de visor no son errores."
    sheet.auto_filter.ref = f"A14:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}"
    note = report.create_sheet("Simulacion")
    note.append(["Motor", "compare_excels — comparador real del programa"])
    for name, count in scenarios.items():
        note.append([name, count])
    note.append(["Controles de visor añadidos (no discrepancias)", controls])
    save_workbook_atomic(report, report_path)
    provenance = json.loads(report_path.with_suffix(".fuentes.json").read_text(encoding="utf-8"))
    write_review_sources(report_path, summary.excel_path, reference_path, list(FIELD_MAP), provenance.get("pdf_sources"))
    result = {"simulation": True, "root": str(output), "excel": summary.excel_path, "reference": str(reference_path), "report": str(report_path), "pdf_total": len(pdfs), "pdf_without_reading": missing, "readings": summary.success, "external_calls": 0, "comparison_findings": stats["hallazgos_total"], "viewer_controls": controls, "scenarios": scenarios}
    write_json_atomic(output / "SIMULACION.json", result)
    return result
