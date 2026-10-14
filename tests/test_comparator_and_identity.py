import pytest
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

from core.comparador import (
    compare_and_update,
    compare_excels,
    discover_pdf_root,
    discrepancy_magnitude,
    parse_date,
    parse_ipress_from_relative_path,
    resolve_source_ipress,
    validate_comparison_inputs,
)


def _reference_csv(tmp_path):
    path = tmp_path / "reference.csv"
    path.write_text("Num Doc.,Fecha de Muestra,HB\n12345678,01/06/2026,10.0\n", encoding="utf-8")
    return str(path)


def test_comparator_handles_missing_document_and_date(tmp_path):
    reference = _reference_csv(tmp_path)
    archivo = [
        {"documento_archivo_pdf": None, "fecha_toma_muestra": "01/06/2026"},
        {"documento_archivo_pdf": "12345678", "fecha_toma_muestra": None},
    ]
    control = [{}, {}]
    sairc = [{}, {}]

    compare_and_update(reference, archivo, control, sairc)

    assert archivo[0]["Discrepancias_Urea_Hemo"] == "Sin documento"
    assert archivo[1]["Discrepancias_Urea_Hemo"] == "Fecha invalida o vacia"


def test_comparator_uses_numeric_tolerance_and_validates_lengths(tmp_path):
    reference = _reference_csv(tmp_path)
    archivo = [{"documento_archivo_pdf": "12345678", "fecha_toma_muestra": "01/06/2026"}]
    control = [{}]
    sairc = [{"HEMOGLOBINA": 10.00000001}]
    compare_and_update(reference, archivo, control, sairc)
    assert archivo[0]["Discrepancias_Urea_Hemo"] == "Coincide"

    with pytest.raises(ValueError, match="misma cantidad"):
        compare_and_update(reference, archivo, [], sairc)




def test_comparison_preflight_and_report_distinguish_missing_values(tmp_path):
    ocr_path = tmp_path / "consolidado.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "SAIRC_Formulario2"
    worksheet.append(["documento_archivo_pdf", "fecha_toma_muestra", "HEMOGLOBINA", "UREA_PRE"])
    worksheet.append(["0012345678", "2026-08-10T12:30:00", 10.0, None])
    worksheet.append([None, "10 de agosto del 2026", 11.0, 90.0])
    workbook.save(ocr_path)

    reference = tmp_path / "sairc.csv"
    reference.write_text(
        "Num Doc.,Fecha de Muestra,HB,UPRE\n12345678,10/08/2026,10,100\n",
        encoding="utf-8",
    )

    preflight = validate_comparison_inputs(str(ocr_path), str(reference))
    assert preflight["filas_extraccion"] == 2
    assert preflight["campos_comparables"] == ["HEMOGLOBINA", "UREA_PRE"]
    assert parse_date("10 de agosto del 2026") == (2026, 8)

    output = tmp_path / "comparado.xlsx"
    stats = compare_excels(str(ocr_path), str(reference), str(output))
    assert stats["sin_documento"] == 1
    assert stats["campos_sin_ocr"] == 1
    assert stats["referencias_duplicadas"] == 0

    final_workbook = load_workbook(output, data_only=False)
    assert final_workbook.sheetnames == ["Discrepancias", "Valores_comparados"]
    result = final_workbook["Discrepancias"]
    headers = [cell.value for cell in result[14]]
    type_column = headers.index("Tipo_hallazgo") + 1
    exam_column = headers.index("Examen") + 1
    findings = {(result.cell(row, type_column).value, result.cell(row, exam_column).value) for row in range(15, result.max_row + 1)}
    assert ("SAIRC_CAMPO_SIN_VALOR_PDF", "UREA_PRE") in findings
    assert ("SIN_DOCUMENTO", "REGISTRO_COMPLETO") in findings
    assert result.cell(row=14, column=type_column).font.bold is True
    assert result.column_dimensions[get_column_letter(type_column)].width == 34
    assert result.freeze_panes == "A15"
    assert result.auto_filter.ref.startswith("A14:")


def test_comparison_preflight_rejects_csv_as_extraction_input(tmp_path):
    (tmp_path / "ocr.csv").write_text("x\n", encoding="utf-8")
    (tmp_path / "reference.csv").write_text("Num Doc.,Fecha de Muestra\n1,01/01/2026\n", encoding="utf-8")
    with pytest.raises(ValueError, match="Excel .xlsx"):
        validate_comparison_inputs(str(tmp_path / "ocr.csv"), str(tmp_path / "reference.csv"))


def test_audit_magnitude_prioritizes_large_urea_and_deprioritizes_small_hematocrit():
    assert discrepancy_magnitude("UREA_PRE", 246.3, 146.3)[2] == "CRITICA"
    assert discrepancy_magnitude("HEMATOCRITO", 35.0, 34.7)[2] == "BAJA"


def test_report_distinguishes_sairc_without_pdf_and_pdf_without_sairc(tmp_path):
    ocr_path = tmp_path / "consolidado.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "SAIRC_Formulario2"
    worksheet.append(["documento_archivo_pdf", "fecha_toma_muestra", "archivo", "HEMOGLOBINA"])
    worksheet.append(["11111111", "10/08/2026", "11111111.pdf", 10.0])
    workbook.save(ocr_path)

    reference = tmp_path / "sairc.csv"
    reference.write_text(
        "Num Doc.,Fecha de Muestra,HB,Nombres y apellidos\n22222222,10/08/2026,9.5,Paciente SAIRC\n",
        encoding="utf-8",
    )
    output = tmp_path / "comparado.xlsx"
    stats = compare_excels(str(ocr_path), str(reference), str(output))
    assert stats["pdf_sin_sairc"] == 1
    assert stats["sairc_sin_pdf"] == 1

    result = load_workbook(output, data_only=False)["Discrepancias"]
    headers = [cell.value for cell in result[14]]
    type_column = headers.index("Tipo_hallazgo") + 1
    types = {result.cell(row, type_column).value for row in range(15, result.max_row + 1)}
    assert {"PDF_SIN_REGISTRO_SAIRC", "SAIRC_SIN_PDF"}.issubset(types)


def test_ipress_is_extracted_from_first_route_folder():
    route = r"16565-CENTRO DE DIALISIS SAN LUIS SAC\443-Lima Este 23-17\DIGITALIZACION_INICIAL\123.pdf"
    assert parse_ipress_from_relative_path(route) == ("16565", "Centro de Dialisis San Luis")
    assert parse_ipress_from_relative_path(r"1566-Tacna 24-4\DIGITALIZACION_INICIAL\123.pdf") == ("", "")


def test_hito4_uses_contracted_ipress_and_prepares_review_pdf(tmp_path):
    relative = Path("16565-CENTRO DE DIALISIS SAN LUIS SAC") / "443-Lima Este 23-17" / "DIGITALIZACION_INICIAL" / "12345678.pdf"
    pdf_root = tmp_path / "pdf_root"
    source_pdf = pdf_root / relative
    source_pdf.parent.mkdir(parents=True)
    source_pdf.write_bytes(b"%PDF-1.4 test")

    ocr_path = tmp_path / "consolidado.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "SAIRC_Formulario2"
    worksheet.append([
        "documento_archivo_pdf", "fecha_toma_muestra", "archivo", "ruta_relativa", "laboratorio", "HEMOGLOBINA"
    ])
    worksheet.append(["12345678", "10/08/2026", source_pdf.name, str(relative), "LABORATORIO EMISOR", 12.0])
    workbook.save(ocr_path)

    reference = tmp_path / "sairc.csv"
    reference.write_text("Num Doc.,Fecha de Muestra,HB\n12345678,10/08/2026,10.0\n", encoding="utf-8")
    output = tmp_path / "HITO_4_VALIDACION" / "reporte.xlsx"
    stale_pdf = output.parent / "PDFS_PARA_REVISION" / "obsoleto.pdf"
    stale_pdf.parent.mkdir(parents=True)
    stale_pdf.write_bytes(b"%PDF obsoleto")
    stats = compare_excels(str(ocr_path), str(reference), str(output), pdf_root=str(pdf_root))

    assert stats["pdf_revision_seleccionados"] == 1
    assert stats["pdf_revision_copiados"] == 1
    # Un PDF preexistente sin registro de propiedad no puede eliminarse como copia propia.
    assert stats["pdf_revision_obsoletos_eliminados"] == 0
    assert stale_pdf.exists()
    assert source_pdf.is_file()
    review_pdf = output.parent / "PDFS_PARA_REVISION" / relative
    assert review_pdf.is_file()

    result = load_workbook(output, data_only=False)["Discrepancias"]
    headers = [cell.value for cell in result[14]]
    row = {header: result.cell(15, index + 1).value for index, header in enumerate(headers)}
    assert row["CUI_IPRESS"] == "16565"
    assert row["Establecimiento"] == "Centro de Dialisis San Luis"
    assert row["Laboratorio_Emisor"] == "LABORATORIO EMISOR"
    assert row["Estado_PDF_Revision"] == "DISPONIBLE_REVISION"
    link_column = headers.index("Ruta_PDF_Revision") + 1
    link_cell = result.cell(15, link_column)
    assert link_cell.value == "Abrir PDF"
    assert link_cell.hyperlink.target == review_pdf.resolve().as_uri()


def test_hito4_discovers_pdf_root_as_sibling_of_results(tmp_path):
    relative = Path("16565-IPRESS") / "443-Contrato" / "DIGITALIZACION_INICIAL" / "123.pdf"
    pdf_root = tmp_path / "Reportes" / "2026-7"
    source = pdf_root / relative
    source.parent.mkdir(parents=True)
    source.write_bytes(b"%PDF")
    excel_path = tmp_path / "Reportes" / "RESULTADOS" / "lote" / "hito2.xlsx"
    excel_path.parent.mkdir(parents=True)
    excel_path.write_bytes(b"")

    assert discover_pdf_root(str(excel_path), [str(relative)]) == pdf_root.resolve()


def test_hito4_recovers_missing_ipress_folder_from_historical_route(tmp_path):
    historical_relative = Path("1566-Tacna 24-4") / "DIGITALIZACION_INICIAL" / "12345678.pdf"
    resolved_relative = Path("25436-CENTRO DEL RIÑON S.A.C") / historical_relative
    pdf_root = tmp_path / "2026-7"
    source_pdf = pdf_root / resolved_relative
    source_pdf.parent.mkdir(parents=True)
    source_pdf.write_bytes(b"%PDF-1.4 test")

    ocr_path = tmp_path / "consolidado.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "SAIRC_Formulario2"
    worksheet.append(["documento_archivo_pdf", "fecha_toma_muestra", "archivo", "ruta_relativa", "HEMOGLOBINA"])
    worksheet.append(["12345678", "10/08/2026", source_pdf.name, str(historical_relative), 12.0])
    workbook.save(ocr_path)

    reference = tmp_path / "sairc.csv"
    reference.write_text("Num Doc.,Fecha de Muestra,HB\n12345678,10/08/2026,10.0\n", encoding="utf-8")
    output = tmp_path / "HITO_4_VALIDACION" / "reporte.xlsx"
    stats = compare_excels(str(ocr_path), str(reference), str(output), pdf_root=str(pdf_root))

    assert stats["pdf_revision_rutas_recuperadas"] == 1
    assert stats["pdf_revision_no_encontrados"] == 0
    assert (output.parent / "PDFS_PARA_REVISION" / resolved_relative).is_file()

    result = load_workbook(output, data_only=False)["Discrepancias"]
    headers = [cell.value for cell in result[14]]
    row = {header: result.cell(15, index + 1).value for index, header in enumerate(headers)}
    assert row["CUI_IPRESS"] == "25436"
    assert row["Establecimiento"] == "Centro del Riñon"
    assert row["Origen_Establecimiento"] == "RUTA_PDF_RECUPERADA"
    assert row["Ruta_relativa"] == str(historical_relative)
    assert "Ruta_relativa_Resuelta" not in headers


def test_hito4_consolidates_serology_and_routine_pdfs_in_same_period(tmp_path):
    ocr_path = tmp_path / "consolidado.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "SAIRC_Formulario2"
    worksheet.append([
        "documento_archivo_pdf", "fecha_toma_muestra", "archivo", "HEMOGLOBINA", "UREA_PRE", "UREA_POST",
        "HIV", "VDRL", "HBSAG", "ANTI_HBS", "ANTI_HBC", "HCV", "HTLV",
    ])
    worksheet.append(["12345678", "16/07/2026", "rutina.pdf", 10.0, 100.0, 30.0, None, None, None, None, None, None, None])
    worksheet.append([
        "12345678", "18/07/2026", "serologia.pdf", None, None, None,
        "NO REACTIVO", "NO REACTIVO", "NO REACTIVO", "REACTIVO", "NO REACTIVO", "NO REACTIVO", "NO REACTIVO",
    ])
    workbook.save(ocr_path)

    reference = tmp_path / "sairc.csv"
    reference.write_text(
        "Num Doc.,Fecha de Muestra,HB,UPRE,UPOST\n12345678,16/07/2026,10,100,30\n",
        encoding="utf-8",
    )
    output = tmp_path / "comparado.xlsx"
    stats = compare_excels(str(ocr_path), str(reference), str(output))

    assert stats["grupos_paciente_periodo"] == 1
    assert stats["grupos_con_multiples_pdf"] == 1
    assert stats["grupos_paneles_separados"] == 1
    assert stats["campos_sin_ocr"] == 0
    assert stats["coinciden"] == 1

    result = load_workbook(output, data_only=False)["Discrepancias"]
    headers = [cell.value for cell in result[14]]
    type_column = headers.index("Tipo_hallazgo") + 1
    types = {
        result.cell(row, type_column).value
        for row in range(15, result.max_row + 1)
        if result.cell(row, type_column).value
    }
    assert "SAIRC_CAMPO_SIN_VALOR_PDF" not in types


def test_hito4_does_not_consolidate_patient_between_different_ipress(tmp_path):
    ocr_path = tmp_path / "consolidado.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "SAIRC_Formulario2"
    worksheet.append([
        "ruta_relativa", "archivo", "documento_archivo_pdf", "fecha_toma_muestra",
        "HEMOGLOBINA", "UREA_PRE", "UREA_POST", "HIV",
    ])
    worksheet.append([
        "10227-IPRESS A/1-Contrato/DIGITALIZACION_INICIAL/serologia.pdf",
        "serologia.pdf", "12345678", "18/07/2026", None, None, None, "NO REACTIVO",
    ])
    worksheet.append([
        "16565-IPRESS B/2-Contrato/DIGITALIZACION_INICIAL/rutina.pdf",
        "rutina.pdf", "12345678", "20/07/2026", 10.0, 100.0, 30.0, None,
    ])
    workbook.save(ocr_path)

    reference = tmp_path / "sairc.csv"
    reference.write_text(
        "Renaes,Num Doc.,Fecha de Muestra,HB,UPRE,UPOST\n"
        "10227,12345678,18/07/2026,11,110,35\n"
        "16565,12345678,20/07/2026,10,100,30\n",
        encoding="utf-8",
    )
    output = tmp_path / "comparado.xlsx"
    stats = compare_excels(str(ocr_path), str(reference), str(output))

    assert stats["correlacion_por_ipress"] is True
    assert stats["grupos_paciente_periodo"] == 2
    assert stats["grupos_con_multiples_pdf"] == 0
    assert stats["grupos_paneles_separados"] == 0
    assert stats["referencias_duplicadas"] == 0
    assert stats["grupos_referencia_ambigua"] == 0
    assert stats["campos_sin_ocr"] == 3

    result = load_workbook(output, data_only=False)["Discrepancias"]
    headers = [cell.value for cell in result[14]]
    rows = [dict(zip(headers, values)) for values in result.iter_rows(min_row=15, values_only=True)]
    missing = [row for row in rows if row["Tipo_hallazgo"] == "SAIRC_CAMPO_SIN_VALOR_PDF"]
    assert {row["CUI_IPRESS"] for row in missing} == {"10227"}
    assert {row["Examen"] for row in missing} == {"HEMOGLOBINA", "UREA_PRE", "UREA_POST"}


def test_hito4_collapses_empty_sairc_result_row_into_one_finding(tmp_path):
    ocr_path = tmp_path / "consolidado.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "SAIRC_Formulario2"
    worksheet.append([
        "ruta_relativa", "archivo", "documento_archivo_pdf", "fecha_toma_muestra",
        "HEMOGLOBINA", "UREA_PRE", "UREA_POST",
    ])
    worksheet.append([
        "10227-IPRESS A/1-Contrato/DIGITALIZACION_INICIAL/resultados.pdf",
        "resultados.pdf", "12345678", "18/07/2026", 10.0, 100.0, 30.0,
    ])
    workbook.save(ocr_path)

    reference = tmp_path / "sairc.csv"
    reference.write_text(
        "Renaes,Num Doc.,Fecha de Muestra,HB,UPRE,UPOST\n"
        "10227,12345678,18/07/2026,,,\n",
        encoding="utf-8",
    )
    output = tmp_path / "comparado.xlsx"
    stats = compare_excels(str(ocr_path), str(reference), str(output))

    assert stats["registros_sairc_sin_resultados"] == 1
    assert stats["campos_sin_sairc"] == 3
    assert stats["hallazgos_total"] == 1

    result = load_workbook(output, data_only=False)["Discrepancias"]
    headers = [cell.value for cell in result[14]]
    rows = [dict(zip(headers, values)) for values in result.iter_rows(min_row=15, values_only=True)]
    assert len(rows) == 1
    assert rows[0]["Tipo_hallazgo"] == "REGISTRO_SAIRC_SIN_RESULTADOS"
    assert rows[0]["Examen"] == "RESULTADOS_COMPLETOS"
    assert rows[0]["Grupo_control"] == "PENALIDAD_HB_UREA"
    assert rows[0]["Prioridad"] == "CRITICA"
    assert "HEMOGLOBINA=10" in rows[0]["Interpretacion_control"]
    assert "UREA_PRE=100" in rows[0]["Interpretacion_control"]
    assert "UREA_POST=30" in rows[0]["Interpretacion_control"]


def test_hito4_recovers_cui_from_standard_pdf_filename():
    cui, name, origin = resolve_source_ipress(
        "1566-Tacna 24-4/DIGITALIZACION_INICIAL/archivo.pdf",
        "12345678_REA_31072026_00025436-31072026094842.pdf",
    )

    assert cui == "25436"
    assert name == ""
    assert origin == "NOMBRE_PDF"


def test_hito4_matches_duplicate_monthly_references_by_exact_date(tmp_path):
    ocr_path = tmp_path / "consolidado.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "SAIRC_Formulario2"
    worksheet.append(["documento_archivo_pdf", "fecha_toma_muestra", "archivo", "HEMOGLOBINA"])
    worksheet.append(["12345678", "10/07/2026", "fecha10.pdf", 10.0])
    worksheet.append(["12345678", "20/07/2026", "fecha20.pdf", 12.0])
    workbook.save(ocr_path)

    reference = tmp_path / "sairc.csv"
    reference.write_text(
        "Num Doc.,Fecha de Muestra,HB\n12345678,10/07/2026,10\n12345678,20/07/2026,12\n",
        encoding="utf-8",
    )
    output = tmp_path / "comparado.xlsx"
    stats = compare_excels(str(ocr_path), str(reference), str(output))

    assert stats["referencias_duplicadas"] == 1
    assert stats["comparaciones_fecha_exacta"] == 2
    assert stats["campos_referencia_ambigua"] == 0
    assert stats["coinciden"] == 1

    result = load_workbook(output, data_only=False)["Discrepancias"]
    headers = [cell.value for cell in result[14]]
    type_column = headers.index("Tipo_hallazgo") + 1
    types = {
        result.cell(row, type_column).value
        for row in range(15, result.max_row + 1)
        if result.cell(row, type_column).value
    }
    assert "VALOR_DISCREPANTE" not in types
    assert "REFERENCIA_DUPLICADA" in types


def test_hito4_handles_folder_with_trailing_spaces(tmp_path):
    # Simulates folder on Windows like '1716-Huaraz 26-1 ' with trailing space
    relative = Path("31041-NEFRO") / "1716-Huaraz 26-1 " / "DIGITALIZACION_INICIAL" / "sample.pdf"
    pdf_root = tmp_path / "Reportes"
    # To create folder with trailing space on Windows without OS error, use _safe_mkdir or standard mkdir
    from core.comparador import _safe_mkdir, _sanitize_path_for_filesystem, _to_win_long_path
    _safe_mkdir(pdf_root / relative.parent)
    source_pdf = pdf_root / relative
    with open(_to_win_long_path(source_pdf), "wb") as f:
        f.write(b"%PDF-test")

    ocr_path = tmp_path / "consolidado.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "SAIRC_Formulario2"
    worksheet.append(["documento_archivo_pdf", "fecha_toma_muestra", "archivo", "ruta_relativa", "HEMOGLOBINA"])
    worksheet.append(["87654321", "10/08/2026", source_pdf.name, str(relative), 15.0])
    workbook.save(ocr_path)

    reference = tmp_path / "sairc.csv"
    reference.write_text("Num Doc.,Fecha de Muestra,HB\n87654321,10/08/2026,10.0\n", encoding="utf-8")
    output = tmp_path / "HITO_4_VALIDACION" / "reporte.xlsx"

    stats = compare_excels(str(ocr_path), str(reference), str(output), pdf_root=str(pdf_root))
    assert stats["pdf_revision_seleccionados"] == 1
    assert stats["pdf_revision_copiados"] == 1

    sanitized_rel = _sanitize_path_for_filesystem(relative)
    assert (output.parent / "PDFS_PARA_REVISION" / sanitized_rel).is_file()


def test_normalize_doc_defensive_filtering():
    from core.sairc_normalizer import normalize_doc

    # Valid Peruvian DNI and documents
    assert normalize_doc("42183251") == "42183251"
    assert normalize_doc(" 00879997 ") == "00879997"
    assert normalize_doc("DNI: 12345678") == "12345678"

    # Explanatory and CUI strings from Mistral or OCR should be rejected to None
    assert normalize_doc("CUI 00023617 (No se identifica DNI/C.E. explícito en el documento)") is None
    assert normalize_doc("CUI 00023617 (Laboratorio SEI Salud)") is None
    assert normalize_doc("No disponible (solo se identifica paciente, no documento)") is None
    assert normalize_doc("No identificado (DNI/C.E. no visible)") is None
    assert normalize_doc("No especificado (no visible en el documento)") is None
    assert normalize_doc("20559272192") is None  # RUC of clinic starting with 20


def test_standardize_date_rejects_times_and_handles_two_digit_years():
    from core.sairc_normalizer import standardize_date

    # Pure time strings should be rejected to None
    assert standardize_date("17:18:06") is None
    assert standardize_date("12:59:39") is None
    assert standardize_date("14:04") is None

    # Valid dates
    assert standardize_date("13/08/2026") == "13/08/2026"
    assert standardize_date("2026-08-13") == "13/08/2026"

    # 2-digit year support
    assert standardize_date("12/08/26") == "12/08/2026"
    assert standardize_date("01-09-26") == "01/09/2026"


def test_parse_date_rejects_times_and_handles_two_digit_years():
    # Pure time strings
    assert parse_date("17:18:06") == (None, None)
    assert parse_date("13:52:26") == (None, None)

    # 4-digit years
    assert parse_date("13/08/2026") == (2026, 8)
    assert parse_date("2026-08-13") == (2026, 8)

    # 2-digit years
    assert parse_date("12/08/26") == (2026, 8)
    assert parse_date("05/09/26") == (2026, 9)


