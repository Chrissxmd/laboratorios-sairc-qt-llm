from openpyxl import load_workbook

from core.exporters import write_excel


def test_sairc_formulario2_contains_the_complete_universe(tmp_path):
    rows = [
        {"documento_identidad": "11111111", "estado_archivo": "OK_CARGABLE", "puede_cargarse_sairc": "SI"},
        {"documento_identidad": "22222222", "estado_archivo": "SAIRC_REVISION", "puede_cargarse_sairc": "NO"},
    ]

    write_excel(tmp_path, [], [], [], rows, [], [], [], [], [])

    workbook = load_workbook(tmp_path / "resultados_laboratorio_sairc.xlsx", data_only=True)
    formulario = workbook["SAIRC_Formulario2"]
    todos = workbook["SAIRC_Todos_Con_Estado"]
    assert formulario.max_row == 3  # encabezado + universo completo
    assert formulario.cell(2, 1).value == "11111111"
    assert formulario.cell(3, 1).value == "22222222"
    assert formulario.freeze_panes == "A2"
    assert formulario["A1"].font.bold is True
    assert formulario.sheet_view.showGridLines is False
    assert todos.max_row == 3  # encabezado + universo completo
