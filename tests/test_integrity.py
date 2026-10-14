"""Regresiones de integridad con datos sintéticos y sin llamadas externas."""
import json
import os
import sqlite3
from pathlib import Path
import pytest
from openpyxl import Workbook, load_workbook
from core import cache_index, comparador, persistence, processor
from core.comparador import prepare_review_pdfs
from core.exporters import write_excel
from core.workspace import find_hito2_excel, get_hito_paths

@pytest.mark.parametrize('relationship', ['same', 'ancestor', 'descendant'])
@pytest.mark.parametrize('operation', ['copy', 'move'])
def test_review_rejects_overlapping_roots_before_modifying_files(tmp_path, relationship, operation):
    source = tmp_path / 'source'
    source.mkdir()
    original = source / 'original.pdf'
    original.write_bytes(b'synthetic original')
    destination = {'same': source, 'ancestor': tmp_path, 'descendant': source / 'review'}[relationship]
    with pytest.raises(ValueError, match='separadas'):
        prepare_review_pdfs([{'Ruta_relativa': 'original.pdf'}], source, destination, operation=operation)
    assert original.read_bytes() == b'synthetic original'
    assert not (source / 'review').exists()

def test_review_only_removes_unchanged_owned_copies_and_checks_contents(tmp_path):
    source, review = (tmp_path / 'source', tmp_path / 'review')
    source.mkdir()
    for name in ('stale.pdf', 'changed.pdf', 'fresh.pdf'):
        (source / name).write_bytes(b'original')
    first = [{'Ruta_relativa': name} for name in ('stale.pdf', 'changed.pdf', 'fresh.pdf')]
    prepare_review_pdfs(first, source, review)
    (review / 'unrelated.pdf').write_bytes(b'unrelated')
    (review / 'changed.pdf').write_bytes(b'modified')
    source_stat = (source / 'fresh.pdf').stat()
    (source / 'fresh.pdf').write_bytes(b'updated!')
    os.utime(source / 'fresh.pdf', ns=(source_stat.st_atime_ns, source_stat.st_mtime_ns))
    stats = prepare_review_pdfs([{'Ruta_relativa': 'fresh.pdf'}], source, review)
    assert stats['obsoletos_eliminados'] == 1
    assert stats['copiados'] == 1
    assert not (review / 'stale.pdf').exists()
    assert (review / 'changed.pdf').read_bytes() == b'modified'
    assert (review / 'unrelated.pdf').read_bytes() == b'unrelated'
    assert (review / 'fresh.pdf').read_bytes() == b'updated!'
    assert all(((source / name).exists() for name in ('stale.pdf', 'changed.pdf', 'fresh.pdf')))

@pytest.mark.parametrize('operation', ['copy', 'move'])
def test_review_does_not_overwrite_a_foreign_file(tmp_path, operation):
    source, review = (tmp_path / 'source', tmp_path / 'review')
    source.mkdir()
    review.mkdir()
    (source / 'a.pdf').write_bytes(b'new')
    (review / 'a.pdf').write_bytes(b'old')
    with pytest.raises(FileExistsError):
        prepare_review_pdfs([{'Ruta_relativa': 'a.pdf'}], source, review, operation=operation)
    assert (source / 'a.pdf').read_bytes() == b'new'
    assert (review / 'a.pdf').read_bytes() == b'old'

def test_excel_failure_preserves_previous_file_and_is_reported(tmp_path, monkeypatch):
    path = tmp_path / 'resultados_laboratorio_sairc.xlsx'
    path.write_bytes(b'previous workbook')

    def fail_save(_workbook, destination):
        Path(destination).write_bytes(b'incomplete workbook')
        raise OSError('synthetic disk failure')
    monkeypatch.setattr(Workbook, 'save', fail_save)
    with pytest.raises(RuntimeError, match='Exportación incompleta'):
        write_excel(tmp_path, [], [], [], [], [], [], [], [], [])
    assert path.read_bytes() == b'previous workbook'
    assert list(tmp_path.iterdir()) == [path]

def test_excel_fallback_returns_actual_path(tmp_path, monkeypatch):
    from core import exporters
    original_save = exporters.save_workbook_atomic
    primary = tmp_path / 'resultados_laboratorio_sairc.xlsx'

    def locked_primary(workbook, path):
        if path == primary:
            raise PermissionError('synthetic locked workbook')
        original_save(workbook, path)
    monkeypatch.setattr(exporters, 'save_workbook_atomic', locked_primary)
    output = write_excel(tmp_path, [], [], [], [], [], [], [], [], [])
    assert output != primary and output.is_file()
    workbook = load_workbook(output)
    assert 'SAIRC_Formulario2' in workbook.sheetnames
    workbook.close()

def test_hito4_selects_actual_export_and_rejects_an_incomplete_run(tmp_path):
    hito2 = get_hito_paths(tmp_path).hito2
    hito2.mkdir()
    primary = hito2 / 'resultados_laboratorio_sairc.xlsx'
    alternative = hito2 / 'resultados_laboratorio_sairc_locked.xlsx'
    primary.write_bytes(b'old workbook')
    alternative.write_bytes(b'new workbook')
    summary_path = processor._resolve_run_summary_path(hito2)
    state = {'estado_ejecucion': 'REPROCESADO_DESDE_CACHE_IA', 'excel_path': str(alternative)}
    persistence.write_json_atomic(summary_path, state)
    assert find_hito2_excel(tmp_path) == alternative
    state['estado_ejecucion'] = 'EXPORTACION_INCOMPLETA'
    persistence.write_json_atomic(summary_path, state)
    assert find_hito2_excel(tmp_path) is None
    assert primary.read_bytes() == b'old workbook'

def test_hito2_records_export_failure_instead_of_success(tmp_path, monkeypatch):
    responses, output = (tmp_path / 'respuestas_ia', tmp_path / 'output')
    responses.mkdir()
    (responses / 'a.json').write_text(json.dumps({'documento_identidad': '12345678', 'resultados': []}), encoding='utf-8')

    def fail_export(*_args):
        raise OSError('synthetic export failure')
    monkeypatch.setattr(processor, '_write_excel', fail_export)
    with pytest.raises(RuntimeError, match='Exportación incompleta del Hito 2'):
        processor.reprocess_ai_cache(str(responses), str(output))
    state = json.loads(processor._resolve_run_summary_path(output).read_text(encoding='utf-8'))
    assert state['estado_ejecucion'] == 'EXPORTACION_INCOMPLETA'

def test_json_replace_failure_keeps_previous_payload_and_removes_temporary(tmp_path, monkeypatch):
    destination = tmp_path / 'state.json'
    destination.write_text('{"previous": true}', encoding='utf-8')

    def fail_replace(*_args):
        raise OSError('synthetic replace failure')
    monkeypatch.setattr(persistence.os, 'replace', fail_replace)
    with pytest.raises(OSError):
        persistence.write_json_atomic(destination, {'new': True})
    assert json.loads(destination.read_text(encoding='utf-8')) == {'previous': True}
    assert list(tmp_path.iterdir()) == [destination]

def test_cache_connections_are_closed_after_reads_and_writes(tmp_path, monkeypatch):
    connections = []
    connect = cache_index._connect

    def track_connect(path):
        connection = connect(path)
        connections.append(connection)
        return connection
    monkeypatch.setattr(cache_index, '_connect', track_connect)
    database = tmp_path / 'cache.sqlite'
    cache_index.upsert_entry(database, {'ruta_relativa': 'a.pdf', 'status': 'success'})
    assert cache_index.get_entry(database, 'a.pdf')['status'] == 'success'
    assert len(cache_index.iter_entries(database)) == 1
    for connection in connections:
        with pytest.raises(sqlite3.ProgrammingError, match='closed'):
            connection.execute('SELECT 1')

def test_pdf_copy_failure_keeps_original_and_previous_copy(tmp_path, monkeypatch):
    source, destination = (tmp_path / 'source.pdf', tmp_path / 'copy.pdf')
    source.write_bytes(b'original')
    destination.write_bytes(b'previous')

    def fail_copy(_source, target):
        Path(target).write_bytes(b'partial')
        raise OSError('synthetic interrupted copy')
    monkeypatch.setattr(comparador.shutil, 'copy2', fail_copy)
    with pytest.raises(OSError):
        comparador._safe_copy2(source, destination)
    assert source.read_bytes() == b'original'
    assert destination.read_bytes() == b'previous'
    assert set(tmp_path.iterdir()) == {source, destination}

def test_sql_multiple_authorizations_do_not_multiply_reference_rows():
    query = (processor.PROJECT_ROOT / 'sql' / 'consulta_sairc.sql').read_text(encoding='utf-8')
    assert 'NOLOCK' not in query
    joins_and_filters = query[query.index('FROM [BDSAIRC_PBI]'):query.index('ORDER BY')]
    joins_and_filters = joins_and_filters.replace('[BDSAIRC_PBI].[dbo].', '')
    connection = sqlite3.connect(':memory:')
    try:
        connection.executescript('\n            CREATE TABLE SAIRT_FormularioEstado (FormularioEstadoId INT, ParametroFormularioId INT, EstablecimientoId INT, Anio INT, Mes INT);\n            CREATE TABLE SAIRT_PARAMETRO (CodigoParametro INT);\n            CREATE TABLE SAIRT_DatoKTVHD (FormularioEstadoId INT, PacienteId INT);\n            CREATE TABLE Paciente (PacienteId INT, PersonaId INT);\n            CREATE TABLE Persona (PersonaId INT);\n            CREATE TABLE Establecimiento (EstablecimientoId INT);\n            CREATE TABLE AutorizacionTratamientoEstablecimiento (PacienteId INT);\n            INSERT INTO SAIRT_FormularioEstado VALUES (1, 2, 3, 2026, 7);\n            INSERT INTO SAIRT_PARAMETRO VALUES (2);\n            INSERT INTO SAIRT_DatoKTVHD VALUES (1, 10), (1, 20);\n            INSERT INTO Paciente VALUES (10, 100), (20, 200);\n            INSERT INTO Persona VALUES (100), (200);\n            INSERT INTO Establecimiento VALUES (3);\n            INSERT INTO AutorizacionTratamientoEstablecimiento VALUES (10), (10);\n        ')
        rows = connection.execute('SELECT ce.PacienteId ' + joins_and_filters, {'Anio': 2026, 'Mes': 7, 'EstablecimientoId': None}).fetchall()
        assert rows == [(10,)]
    finally:
        connection.close()
