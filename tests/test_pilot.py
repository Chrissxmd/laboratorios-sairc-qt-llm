import threading
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from core import ai, pilot
from core.ai import AIConfig, AIError, MISTRAL_PILOT_MODELS


def _config():
    return AIConfig(provider="mistral", model=MISTRAL_PILOT_MODELS[0], reader_mode="vision", api_key="synthetic", allow_external_processing=True)


def test_pilot_separates_models_reuses_cache_and_locks_dataset(make_pdf, tmp_path, monkeypatch):
    source = tmp_path / "pdfs"
    source.mkdir()
    (source / "synthetic.pdf").write_bytes(make_pdf().read_bytes())
    calls = []
    catalog = {"data": [{"id": model, "capabilities": {"vision": True}} for model in MISTRAL_PILOT_MODELS]}
    monkeypatch.setattr(pilot, "_get", lambda *a, **kw: catalog)
    monkeypatch.setattr(ai, "_get", lambda *a, **kw: catalog)
    def read(config, **kwargs):
        calls.append(config.model)
        value = "13.0" if config.model == MISTRAL_PILOT_MODELS[2] else "12.0"
        result = ai.validate_lab_json({"v": 2, "d": "12345678", "ft": "01/06/2026", "r": [{"a": "HEMOGLOBINA", "v": value, "u": "g/dL"}]})
        result["_meta"]["token_usage"] = {"prompt_token_count": 10000, "candidates_token_count": 2000, "total_token_count": 12000, "cost_estimated_usd": ai.estimate_cost_usd("mistral", config.model, 10000, 2000)}
        return result
    monkeypatch.setattr(ai, "analyze_to_json", read)
    workspace = tmp_path / "results"
    first = pilot.run_mistral_pilot(source, workspace, _config(), authorize_three_models=True)
    assert calls == list(MISTRAL_PILOT_MODELS)
    assert [row["costo_acumulado_usd"] for row in first["models"]] == [0.0027, 0.0018, 0.0012]
    assert all(row["exitos"] == 1 and row["excel_sairc"] for row in first["models"])
    workbook = load_workbook(first["report"], read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["Resumen", "Lecturas_comparadas", "Interpretacion"]
        rows = list(workbook["Lecturas_comparadas"].values)
        hb = next(row for row in rows if row[1] == "HEMOGLOBINA")
        assert hb[2:5] == ("12", "12", "13") and hb[5] == "NO"
    finally:
        workbook.close()
    resumed = pilot.run_mistral_pilot(source, workspace, _config(), authorize_three_models=True)
    assert len(calls) == 3 and all(row["reutilizados"] == 1 for row in resumed["models"])
    model_root = Path(first["output_dir"]) / MISTRAL_PILOT_MODELS[0]
    state_path = next(model_root.rglob("*.meta.json"))
    state = json.loads(state_path.read_text(encoding="utf-8"))
    state["status"] = "error"
    state_path.write_text(json.dumps(state), encoding="utf-8")
    # Un JSON anterior no debe entrar al Excel del piloto si el estado actual falló.
    filtered = pilot.reprocess_ai_cache(str(model_root / "HITO_1_EXTRACCION"), str(tmp_path / "filtered"), state_dir=str(state_path.parent))
    assert filtered.total == 0
    (source / "synthetic.pdf").write_bytes(b"changed synthetic")
    with pytest.raises(ValueError, match="conjunto de PDF cambió"):
        pilot.run_mistral_pilot(source, workspace, _config(), authorize_three_models=True)
    assert len(calls) == 3


def test_pilot_rejects_missing_model_and_missing_consent_before_generation(make_pdf, tmp_path, monkeypatch):
    source = tmp_path / "pdfs"
    source.mkdir()
    (source / "synthetic.pdf").write_bytes(make_pdf().read_bytes())
    monkeypatch.setattr(ai, "analyze_to_json", lambda *a, **kw: pytest.fail("No debe enviar documentos"))
    monkeypatch.setattr(pilot, "_get", lambda *a, **kw: {"data": [{"id": model} for model in MISTRAL_PILOT_MODELS[:2]]})
    with pytest.raises(ValueError, match="autoriza"):
        pilot.run_mistral_pilot(source, tmp_path / "results", _config())
    with pytest.raises(AIError, match="no está disponible"):
        pilot.run_mistral_pilot(source, tmp_path / "results", _config(), authorize_three_models=True)
    assert not (tmp_path / "results").exists()


def test_pilot_cancel_between_models_keeps_partial_report(make_pdf, tmp_path, monkeypatch):
    source = tmp_path / "pdfs"
    source.mkdir()
    (source / "synthetic.pdf").write_bytes(make_pdf().read_bytes())
    catalog = {"data": [{"id": model} for model in MISTRAL_PILOT_MODELS]}
    monkeypatch.setattr(pilot, "_get", lambda *a, **kw: catalog)
    monkeypatch.setattr(ai, "_get", lambda *a, **kw: catalog)
    monkeypatch.setattr(ai, "analyze_to_json", lambda *a, **kw: ai.validate_lab_json({"v": 2, "r": []}))
    cancel = threading.Event()
    original = pilot.process_folder
    def extract(*args, **kwargs):
        result = original(*args, **kwargs)
        cancel.set()
        return result
    monkeypatch.setattr(pilot, "process_folder", extract)
    result = pilot.run_mistral_pilot(source, tmp_path / "results", _config(), authorize_three_models=True, cancel_event=cancel)
    assert result["cancelled"] and len(result["models"]) == 1 and result["report"]
