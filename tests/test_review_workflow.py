import copy
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from core.comparador import compare_excels, validate_comparison_inputs
from core.review_state import ReviewStore, prepare_findings
from ui_qt.review import ReviewPane, read_discrepancies


def finding(pdf, field="HEMOGLOBINA", document="12345678", priority="ALTA"):
    return {"Documento": document, "Fecha_periodo": "2026-06-01", "Ruta_relativa": pdf.name, "_pdf_path": str(pdf), "Examen": field, "Tipo_hallazgo": "VALOR_DISCREPANTE", "Prioridad": priority, "Valor_LLM_OCR": 12.0, "Valor_SAIRC": 120.0}


def report_file(tmp_path, rows):
    book = Workbook()
    sheet = book.active
    sheet.title = "Discrepancias"
    headers = ["Documento", "Fecha_periodo", "Ruta_relativa", "Ruta_PDF_Revision", "Examen", "Tipo_hallazgo", "Prioridad", "Valor_LLM_OCR", "Valor_SAIRC"]
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(key) for key in headers])
        sheet.cell(sheet.max_row, 4).hyperlink = Path(row["_pdf_path"]).as_uri()
    path = tmp_path / "report.xlsx"
    book.save(path)
    return path


def test_pdf_group_decisions_next_and_page_preserved(make_pdf, qt_app, tmp_path):
    first = make_pdf("first.pdf", pages=3)
    second = make_pdf("second.pdf")
    report = report_file(tmp_path, [finding(first), finding(first, "UREA_PRE"), finding(second, priority="BAJA")])
    pane = ReviewPane()
    pane.set_rows(read_discrepancies(report))
    assert pane.source.rowCount() == 2 and pane.findings_source.rowCount() == 2
    pane.viewer.page.setValue(2)
    pane.viewer.zoom.setCurrentIndex(3)
    pane.viewer.search.setText("HEMOGLOBINA")
    pane.findings.selectRow(1)
    assert pane.viewer.page.value() == 2 and pane.viewer.zoom.currentIndex() == 3
    assert pane.viewer.search.text() == "HEMOGLOBINA"
    pane.decision.setCurrentText("Confirma Hito 2")
    pane.save_decision(True)
    assert pane.current_case["Estado"] == "En revisión"
    assert pane.viewer.path == first.resolve()
    pane.decision.setCurrentText("Confirma Hito 3")
    pane.save_decision(True)
    assert pane.viewer.path == second.resolve()
    assert pane.source.rows[0]["Estado"] == "Completado"
    pane.decision.setCurrentText("No concluyente")
    pane.reason.setText("Lectura no concluyente sintética")
    pane.save_decision()
    assert pane.current_case["Estado"] == "Derivado"
    assert "1 completados" in pane.count.text() and "1 derivados" in pane.count.text()
    pane.undo()
    assert pane.current_case["Estado"] == "Pendiente"
    pane.close()


def test_resume_decisions_filter_field_and_viewer_position(make_pdf, qt_app, tmp_path):
    pdf = make_pdf(pages=3)
    report = report_file(tmp_path, [finding(pdf), finding(pdf, "UREA_PRE")])
    pane = ReviewPane()
    pane.set_rows(read_discrepancies(report))
    pane.state.setCurrentText("Todos")
    pane.query.setText("12345678")
    pane.findings.selectRow(1)
    pane.viewer.page.setValue(3)
    pane.viewer.zoom.setCurrentIndex(4)
    pane.decision.setCurrentText("Confirma Hito 2")
    pane.save_decision()
    pane.close()
    other = ReviewPane()
    other.set_rows(read_discrepancies(report))
    assert other.state.currentText() == "Todos" and other.query.text() == "12345678"
    assert other.findings.currentIndex().row() == 1
    assert other.viewer.page.value() == 3 and other.viewer.zoom.currentIndex() == 4
    assert len(other.store.decisions) == 1
    other.undo()
    assert not other.store.decisions
    history = other.store.connection.execute("SELECT action,reviewer,updated FROM history ORDER BY id").fetchall()
    assert history[-1][0] == "DESHACER" and all(row[1] and row[2] for row in history)
    other.close()


def test_changed_pdf_invalidates_only_its_decisions(make_pdf, tmp_path):
    first, second = make_pdf("a.pdf"), make_pdf("b.pdf")
    report = report_file(tmp_path, [finding(first), finding(second)])
    rows = read_discrepancies(report)
    store = ReviewStore(rows, report)
    for row in rows:
        store.save(row["_id"], "Confirma Hito 2", page=1)
    first.write_bytes(first.read_bytes() + b"\n% synthetic change\n")
    with pytest.raises(ValueError, match="cambió"):
        store.export(tmp_path / "stale.xlsx")
    store.close()
    changed = read_discrepancies(report)
    other = ReviewStore(changed, report)
    assert len(other.decisions) == 1
    assert changed[1]["_id"] in other.decisions
    assert changed[0]["_id"] not in other.decisions
    other.close()


def test_reference_change_requires_new_comparison(make_pdf, tmp_path):
    pdf = make_pdf()
    original_root = tmp_path / "originals"
    original_root.mkdir()
    original_pdf = original_root / pdf.name
    original_pdf.write_bytes(pdf.read_bytes())
    pdf = original_pdf
    excel = tmp_path / "hito2.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = "SAIRC_Formulario2"
    sheet.append(["documento_archivo_pdf", "fecha_toma_muestra", "ruta_relativa", "archivo", "HEMOGLOBINA", "UREA_PRE"])
    sheet.append(["12345678", "01/06/2026", pdf.name, pdf.name, 12, 80])
    book.save(excel)
    reference = tmp_path / "ref.csv"
    reference.write_text("Num Doc.,Fecha de Muestra,HB,UPRE\n12345678,01/06/2026,120,80\n")
    report = tmp_path / "comparison.xlsx"
    compare_excels(str(excel), str(reference), str(report), pdf_root=str(original_root))
    rows = read_discrepancies(report)
    assert rows[0]["_context"] and any(item["Estado"] == "COINCIDE" for item in rows[0]["_context"])
    store = ReviewStore(rows, report)
    store.save(rows[0]["_id"], "Confirma Hito 2", page=1)
    store.close()
    original_bytes = pdf.read_bytes()
    pdf.write_bytes(original_bytes + b"\n% changed original, unchanged review copy\n")
    modified_pdf = read_discrepancies(report)
    changed_pdf_store = ReviewStore(modified_pdf, report)
    assert not changed_pdf_store.decisions
    assert "PDF original" in modified_pdf[0]["_warning"]
    with pytest.raises(ValueError, match="PDF original"):
        changed_pdf_store.save(modified_pdf[0]["_id"], "Confirma Hito 2", page=1)
    changed_pdf_store.close()
    pdf.write_bytes(original_bytes)
    reference.write_text("Num Doc.,Fecha de Muestra,HB,UPRE\n12345678,01/06/2026,12,80\n")
    changed = read_discrepancies(report)
    other = ReviewStore(changed, report)
    assert not other.decisions and changed[0]["_warning"]
    with pytest.raises(ValueError, match="fuente"):
        other.save(changed[0]["_id"], "Confirma Hito 2", page=1)
    other.close()


def test_write_only_excel_without_dimensions_can_be_compared(tmp_path):
    book = Workbook(write_only=True)
    sheet = book.create_sheet("SAIRC_Formulario2")
    sheet.append(["documento_archivo_pdf", "fecha_toma_muestra", "HEMOGLOBINA"])
    sheet.append(["12345678", "01/06/2026", 12])
    excel = tmp_path / "streamed.xlsx"
    book.save(excel)
    reference = tmp_path / "ref.csv"
    reference.write_text("Num Doc.,Fecha de Muestra,HB\n12345678,01/06/2026,13\n")
    assert validate_comparison_inputs(str(excel), str(reference))["filas_extraccion"] == 1
    result = compare_excels(str(excel), str(reference), str(tmp_path / "report.xlsx"))
    assert result["filas_ocr"] == 1 and result["hallazgos_total"] == 1


def test_export_proposals_original_values_no_formulas_or_source_changes(make_pdf, tmp_path):
    pdf = make_pdf()
    report = report_file(tmp_path, [finding(pdf)])
    originals = {path: path.read_bytes() for path in (pdf, report)}
    rows = read_discrepancies(report)
    store = ReviewStore(rows, report)
    with pytest.raises(ValueError, match="motivo"):
        store.save(rows[0]["_id"], "No concluyente")
    with pytest.raises(ValueError, match="página"):
        store.save(rows[0]["_id"], "Confirma Hito 2")
    store.save(rows[0]["_id"], "Otro valor del PDF", value="=1+1", reason="Texto sintético de prueba", page=1)
    output = tmp_path / "proposals.xlsx"
    store.export(output)
    book = load_workbook(output)
    sheet = book.active
    assert sheet.cell(2, 6).value == 12 and sheet.cell(2, 7).value == 120
    assert sheet.cell(2, 9).value == "=1+1" and sheet.cell(2, 9).data_type == "s"
    assert sheet.cell(2, 15).value == "PROPUESTA_NO_APLICADA"
    book.close()
    assert all(path.read_bytes() == value for path, value in originals.items())
    store.close()


def test_missing_pdf_cannot_be_confirmed_but_can_be_derived(tmp_path):
    rows = prepare_findings([finding(tmp_path / "missing.pdf")])
    store = ReviewStore(rows)
    with pytest.raises(ValueError, match="PDF disponible"):
        store.save(rows[0]["_id"], "Confirma Hito 3", page=1)
    store.save(rows[0]["_id"], "No concluyente", reason="PDF no disponible")
    assert store.decisions[rows[0]["_id"]]["choice"] == "No concluyente"
    store.close()


def test_queue_priorities_not_alphabetic_and_filter_orphans(make_pdf, qt_app):
    pdf = make_pdf()
    rows = [finding(pdf, priority="BAJA"), {**finding(pdf, document="99999999", priority="CRITICA"), "Ruta_relativa": "", "_pdf_path": "", "Tipo_hallazgo": "SAIRC_SIN_PDF"}]
    pane = ReviewPane()
    pane.set_rows(copy.deepcopy(rows))
    assert pane.current_case["Prioridad"] == "CRITICA"
    assert pane.viewer.path is None
    pane.kind.setCurrentText("SAIRC_SIN_PDF")
    assert pane.proxy.rowCount() == 1
    pane.close()
