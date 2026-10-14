"""Carga sintética reproducible: 8.000 registros, 2.000 casos y 26.000 hallazgos.

No OCR, IA, SQL ni PDF reales. Mide comparación/importación/cola; no rasterizado.
"""
import ctypes
import argparse
import json
import os
import sys
import time
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook
from PySide6.QtWidgets import QApplication
from core.comparador import FIELD_MAP, compare_excels
from ui_qt.review import ReviewPane, read_discrepancies


def memory_mb():
    class Counters(ctypes.Structure):
        _fields_ = [("cb", ctypes.c_ulong), ("PageFaultCount", ctypes.c_ulong), *[(name, ctypes.c_size_t) for name in ("PeakWorkingSetSize", "WorkingSetSize", "QuotaPeakPagedPoolUsage", "QuotaPagedPoolUsage", "QuotaPeakNonPagedPoolUsage", "QuotaNonPagedPoolUsage", "PagefileUsage", "PeakPagefileUsage")]]
    counters = Counters()
    counters.cb = ctypes.sizeof(counters)
    handle = ctypes.windll.kernel32.GetCurrentProcess
    handle.restype = ctypes.c_void_p
    read = ctypes.windll.psapi.GetProcessMemoryInfo
    read.argtypes = [ctypes.c_void_p, ctypes.POINTER(Counters), ctypes.c_ulong]
    read(handle(), ctypes.byref(counters), counters.cb)
    return round(counters.PeakWorkingSetSize / 1024**2, 1)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--records", type=int, default=8000)
    parser.add_argument("--cases", type=int, default=2000)
    args = parser.parse_args()
    if not 0 < args.cases <= args.records:
        parser.error("Se requiere 0 < cases <= records")
    app = QApplication([])
    root = Path(f"tmp/carga_revision_{args.records}").resolve()
    root.mkdir(parents=True, exist_ok=True)
    originals = root / "originals_vacios"
    originals.mkdir(exist_ok=True)
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("SAIRC_Formulario2")
    sheet.append(["documento_archivo_pdf", "fecha_toma_muestra", "ruta_relativa", "archivo", *FIELD_MAP])
    ref_book = Workbook(write_only=True)
    ref = ref_book.create_sheet("Referencia")
    ref.append(["Num Doc.", "Fecha de Muestra", *FIELD_MAP.values()])
    for i in range(args.records):
        document = str(40000000+i)
        values = [float(20+j) for j in range(len(FIELD_MAP))]
        sheet.append([document, "2026-06-01", f"{document}.pdf", f"{document}.pdf", *values])
        ref.append([document, "2026-06-01", *([value*10 for value in values] if i<args.cases else values)])
    excel, reference, report = root / "hito2.xlsx", root / "hito3.xlsx", root / "hito4.xlsx"
    workbook.save(excel)
    ref_book.save(reference)
    started = time.perf_counter()
    stats = compare_excels(str(excel), str(reference), str(report), pdf_root=str(originals))
    comparison = time.perf_counter()-started
    started = time.perf_counter()
    rows = read_discrepancies(report)
    loading = time.perf_counter()-started
    pane = ReviewPane()
    started = time.perf_counter()
    pane.set_rows(rows)
    app.processEvents()
    grouping = time.perf_counter()-started
    assert len(rows)==args.cases*len(FIELD_MAP) and pane.source.rowCount()==args.cases
    started = time.perf_counter()
    pane.query.setText("40000001")
    app.processEvents()
    filtering = time.perf_counter()-started
    assert pane.proxy.rowCount()==1
    pane.query.clear()
    started = time.perf_counter()
    for _ in range(50):
        pane.step(1)
        app.processEvents()
    navigation = (time.perf_counter()-started)/50
    result = {"scope":"SINTETICO: registros y cola, sin PDF renderizados ni IA/SQL", "records":args.records,"cases":args.cases,"findings":len(rows),"comparison_seconds":round(comparison,3),"report_load_seconds":round(loading,3),"queue_seconds":round(grouping,3),"filter_ms":round(filtering*1000,2),"navigation_mean_ms":round(navigation*1000,2),"peak_process_memory_mb":memory_mb(),"equal_groups":stats["coinciden"]}
    pane.close()
    (root/"medicion.json").write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(result,ensure_ascii=False,indent=2))


if __name__ == "__main__":
    main()
