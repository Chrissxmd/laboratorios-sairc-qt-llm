from html import escape
from pathlib import Path
import posixpath
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from PySide6.QtCore import QAbstractTableModel, QModelIndex, QSortFilterProxyModel, Qt, QUrl
from PySide6.QtGui import QKeySequence, QShortcut
from PySide6.QtWidgets import QAbstractItemView, QComboBox, QDialog, QDialogButtonBox, QFileDialog, QFormLayout, QHBoxLayout, QHeaderView, QLabel, QLineEdit, QPushButton, QScrollArea, QSplitter, QTableView, QTextBrowser, QVBoxLayout, QWidget

from core.review_state import CHOICES, PRIORITIES, RESOLVED, ReviewStore, case_key, prepare_findings
from ui_qt.pdf_viewer import PdfViewer


def report_links(path):
    """Hipervínculos en flujo: openpyxl read_only no expone cell.hyperlink."""
    rel = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    main = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(path) as archive:
        book = ElementTree.fromstring(archive.read("xl/workbook.xml"))
        identity = next(sheet.attrib[rel] for sheet in book.iter(main + "sheet") if sheet.attrib.get("name") == "Discrepancias")
        relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        target = next(item.attrib["Target"] for item in relationships if item.attrib.get("Id") == identity)
        sheet_path = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
        links_path = posixpath.join(posixpath.dirname(sheet_path), "_rels", posixpath.basename(sheet_path) + ".rels")
        if links_path not in archive.namelist():
            return {}
        targets = {item.attrib["Id"]: item.attrib.get("Target", "") for item in ElementTree.fromstring(archive.read(links_path))}
        links = {}
        with archive.open(sheet_path) as stream:
            for _event, element in ElementTree.iterparse(stream, events=("end",)):
                if element.tag == main + "hyperlink":
                    links[element.attrib.get("ref")] = targets.get(element.attrib.get(rel), "")
                element.clear()
        return links


def read_discrepancies(path, pdf_root=""):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Discrepancias" not in workbook.sheetnames:
            raise ValueError("El reporte debe contener la hoja Discrepancias.")
        sheet = workbook["Discrepancias"]
        simulation = "Simulacion" in workbook.sheetnames
        iterator = sheet.iter_rows(values_only=True)
        for header_row, values_row in enumerate(iterator, start=1):
            if header_row == 1:
                simulation = simulation or "SIMULACIÓN" in str(values_row[0] or "").upper()
            if header_row > 30:
                raise ValueError("No se encontraron encabezados de discrepancias válidos.")
            headers = [str(value or "").strip() for value in values_row]
            if {"Tipo_hallazgo", "Documento", "Ruta_relativa"}.issubset(headers):
                break
        else:
            raise ValueError("No se encontraron encabezados de discrepancias válidos.")
        rows = []
        links = report_links(path)
        for number, cells in enumerate(iterator, start=header_row + 1):
            row = {header: value for header, value in zip(headers, cells) if header}
            if not row.get("Tipo_hallazgo"):
                continue
            row["_pdf_path"] = ""
            row["_simulation"] = simulation or str(row["Tipo_hallazgo"]).startswith("SIMULACION_")
            for index, header in enumerate(headers, start=1):
                if header == "Ruta_PDF_Revision":
                    url = QUrl(links.get(f"{get_column_letter(index)}{number}", ""))
                    if url.isLocalFile():
                        row["_pdf_path"] = url.toLocalFile()
            if pdf_root and not Path(row["_pdf_path"]).is_file():
                root = Path(pdf_root).resolve()
                relative = str(row.get("Ruta_relativa") or "").replace("\\", "/")
                if relative and not Path(relative).is_absolute():
                    candidate = (root / relative).resolve()
                    if candidate.is_relative_to(root) and candidate.is_file():
                        row["_pdf_path"] = str(candidate)
            rows.append(row)
        context = {}
        if "Valores_comparados" in workbook.sheetnames:
            values = workbook["Valores_comparados"].iter_rows(values_only=True)
            context_headers = next(values)
            for values_row in values:
                item = dict(zip(context_headers, values_row))
                key = str(item.get("Ruta_relativa") or "").replace("\\", "/").casefold()
                context.setdefault(key, []).append(item)
        for row in rows:
            row["_context"] = context.get(str(row.get("Ruta_relativa") or "").replace("\\", "/").casefold(), [])
    finally:
        workbook.close()
    return prepare_findings(rows, path)


class FindingsModel(QAbstractTableModel):
    columns = ["Prioridad", "Estado", "PDF", "Documento", "CUI_IPRESS", "Diferencias"]

    def __init__(self, rows, parent=None, columns=None):
        super().__init__(parent)
        self.rows = rows
        self.columns = columns or self.columns

    def rowCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.rows)

    def columnCount(self, parent=QModelIndex()):
        return 0 if parent.isValid() else len(self.columns)

    def data(self, index, role=Qt.ItemDataRole.DisplayRole):
        if not index.isValid():
            return None
        key = self.columns[index.column()]
        value = self.rows[index.row()].get(key)
        if role in (Qt.ItemDataRole.DisplayRole, Qt.ItemDataRole.ToolTipRole):
            return "" if value is None else str(value)
        if role == Qt.ItemDataRole.UserRole:
            if key == "Prioridad":
                return PRIORITIES.get(value, 9)
            return value if isinstance(value, (int, float)) else str(value or "")
        return None

    def headerData(self, section, orientation, role=Qt.ItemDataRole.DisplayRole):
        if role == Qt.ItemDataRole.DisplayRole and orientation == Qt.Orientation.Horizontal:
            return self.columns[section].replace("_", " ")
        return super().headerData(section, orientation, role)


class FindingsFilter(QSortFilterProxyModel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.query, self.priority, self.kind, self.state, self.ipress = "", "Todas", "Todos", "Pendientes", "Todas"
        self.setSortRole(Qt.ItemDataRole.UserRole)
        self.setDynamicSortFilter(False)  # Guardar no reordena ni hace saltar la fila activa.

    def filterAcceptsRow(self, row, parent):
        data = self.sourceModel().rows[row]
        state = self.state == "Todos" or (self.state == "Pendientes" and data["Estado"] in {"Pendiente", "En revisión"}) or data["Estado"] == self.state
        return state and (self.priority == "Todas" or data.get("Prioridad") == self.priority) and (self.kind == "Todos" or self.kind in data["_kinds"]) and (self.ipress == "Todas" or data.get("_ipress") == self.ipress) and self.query.casefold() in data["_search"]


class ReviewPane(QWidget):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.pdf_root, self.rows, self.store, self.current_case = "", [], None, None
        self.restoring = False
        self.source = FindingsModel([], self)
        self.proxy = FindingsFilter(self)
        self.proxy.setSourceModel(self.source)
        self.query = QLineEdit()
        self.query.setPlaceholderText("Buscar documento, IPRESS o examen")
        self.priority, self.kind, self.state, self.ipress = QComboBox(), QComboBox(), QComboBox(), QComboBox()
        self.priority.addItems(["Todas", "CRITICA", "ALTA", "MEDIA", "BAJA"])
        self.kind.addItem("Todos")
        self.ipress.addItem("Todas")
        self.state.addItems(["Pendientes", "Todos", "En revisión", "Completado", "Derivado"])
        self.count = QLabel("Sin reporte cargado")
        self.count.setWordWrap(True)
        self.table = self.make_table(self.proxy)
        self.table.setMaximumHeight(115)
        self.table.setMinimumHeight(80)
        self.table.setSortingEnabled(True)
        self.table.selectionModel().currentRowChanged.connect(self.selected)
        self.findings_source = FindingsModel([], self, ["Examen", "Hito 2", "Hito 3", "Decisión"])
        self.findings = self.make_table(self.findings_source)
        self.findings.setMinimumHeight(100)
        self.findings.setMaximumHeight(170)
        self.findings.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.findings.selectionModel().currentRowChanged.connect(self.field_selected)
        self.details = QTextBrowser()
        self.details.setMaximumHeight(75)
        self.details.setMinimumHeight(45)
        self.decision = QComboBox()
        self.decision.addItems(["Seleccionar decisión…", *CHOICES, "Reabrir"])
        self.manual_value = QLineEdit()
        self.manual_value.setPlaceholderText("Valor confirmado en el PDF")
        self.manual_value.hide()
        self.reason = QLineEdit()
        self.reason.setPlaceholderText("Motivo / observación (obligatorio para otro valor o derivar)")
        self.feedback = QLabel("Las decisiones se guardan aparte; no actualizan SQL ni las fuentes.")
        self.feedback.setWordWrap(True)
        self.save_button = QPushButton("Guardar decisión")
        self.save_button.clicked.connect(self.save_decision)
        self.next_button = QPushButton("Guardar y seguir")
        self.next_button.clicked.connect(lambda: self.save_decision(True))
        self.undo_button = QPushButton("Deshacer")
        self.undo_button.clicked.connect(self.undo)
        self.export_button = QPushButton("Propuestas…")
        self.export_button.setToolTip("Exportar las decisiones como propuestas, sin aplicar a las fuentes")
        self.export_button.clicked.connect(self.export_decisions)
        self.viewer = PdfViewer()
        left = QWidget()
        left_layout = QVBoxLayout(left)
        left_layout.setContentsMargins(0, 0, 0, 0)
        left_layout.addWidget(self.query)
        filters = QHBoxLayout()
        filters.addWidget(self.state, 1)
        filters.addWidget(self.priority, 1)
        filters_button = QPushButton("Filtros…")
        filters_button.clicked.connect(self.filter_dialog)
        filters.addWidget(filters_button)
        left_layout.addLayout(filters)
        left_layout.addWidget(self.count)
        left_layout.addWidget(self.table)
        comparison_label = QLabel("Diferencias del PDF · Hito 2: lectura IA / Hito 3: SAIRC")
        comparison_label.setWordWrap(True)
        self.details.hide()
        controls = QWidget()
        controls_layout = QVBoxLayout(controls)
        controls_layout.setContentsMargins(0, 0, 0, 0)
        controls_layout.addWidget(self.decision)
        controls_layout.addWidget(self.manual_value)
        controls_layout.addWidget(self.reason)
        actions = QHBoxLayout()
        for button in (self.save_button, self.next_button):
            actions.addWidget(button)
        controls_layout.addLayout(actions)
        extra = QHBoxLayout()
        extra.addWidget(self.undo_button)
        extra.addWidget(self.export_button)
        detail_button = QPushButton("Detalles…")
        detail_button.setToolTip("Ver coincidencias, cobertura y datos de correlación")
        detail_button.clicked.connect(self.show_details)
        extra.addWidget(detail_button)
        controls_layout.addLayout(extra)
        controls_layout.addWidget(self.feedback)
        navigation = QHBoxLayout()
        for label, offset in (("Anterior PDF", -1), ("Siguiente PDF", 1)):
            button = QPushButton(label)
            button.clicked.connect(lambda _checked=False, delta=offset: self.step(delta))
            navigation.addWidget(button)
        controls_layout.addLayout(navigation)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setWidget(left)
        scroll.setMinimumHeight(40)
        scroll.setMaximumHeight(250)
        panel = QWidget()
        panel.setMinimumWidth(300)
        panel_layout = QVBoxLayout(panel)
        panel_layout.setContentsMargins(0, 0, 0, 0)
        panel_layout.addWidget(scroll, 1)
        panel_layout.addWidget(comparison_label)
        panel_layout.addWidget(self.findings, 1)
        panel_layout.addWidget(controls)
        self.viewer.setMinimumWidth(480)
        split = QSplitter()
        self.splitter = split
        split.setChildrenCollapsible(False)
        split.addWidget(panel)
        split.addWidget(self.viewer)
        split.setStretchFactor(0, 2)
        split.setStretchFactor(1, 3)
        split.setSizes([480, 720])
        layout = QVBoxLayout(self)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addWidget(split)
        self.query.textChanged.connect(self.apply_filters)
        for field in (self.priority, self.kind, self.state, self.ipress):
            field.currentTextChanged.connect(self.apply_filters)
        self.decision.currentTextChanged.connect(lambda value: self.manual_value.setVisible(value == "Otro valor del PDF"))
        for key, callback in (("Alt+Right", lambda: self.step(1)), ("Alt+Left", lambda: self.step(-1)), ("Alt+G", self.save_decision), ("Alt+N", self.next_pending)):
            shortcut = QShortcut(QKeySequence(key), self)
            shortcut.activated.connect(callback)
        self.update_actions()

    @staticmethod
    def make_table(model):
        table = QTableView()
        table.setModel(model)
        table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        table.setAlternatingRowColors(True)
        table.verticalHeader().setDefaultSectionSize(35)
        return table

    def set_rows(self, rows, report_path=None):
        self.save_session()
        if self.store:
            self.store.close()
        self.store = None
        self.current_case = None
        self.rows = rows if not rows or "_id" in rows[0] else prepare_findings(rows, report_path)
        report = report_path or (rows[0].get("_report") if rows else None)
        self.feedback.setText("Las decisiones se guardan aparte; no actualizan SQL ni las fuentes.")
        try:
            self.store = ReviewStore(self.rows, report)
        except Exception as exc:
            self.feedback.setText(f"No se puede guardar revisión: {exc}. Visor disponible en solo lectura.")
        grouped = {}
        for row in self.rows:
            key = case_key(row)
            if key not in grouped:
                grouped[key] = {"_key": key, "_findings": [], "Prioridad": row.get("Prioridad", "BAJA"), "Documento": row.get("Documento"), "CUI_IPRESS": row.get("CUI_IPRESS"), "PDF": Path(str(row.get("Ruta_relativa") or row.get("_pdf_path") or "Sin PDF")).name, "_ipress": str(row.get("CUI_IPRESS") or row.get("Establecimiento") or "Sin IPRESS")}
            case = grouped[key]
            case["_findings"].append(row)
            if PRIORITIES.get(row.get("Prioridad"), 9) < PRIORITIES.get(case["Prioridad"], 9):
                case["Prioridad"] = row["Prioridad"]
        cases = list(grouped.values())
        for case in cases:
            case["Diferencias"] = len(case["_findings"])
            case["_kinds"] = {row.get("Tipo_hallazgo") for row in case["_findings"]}
            case["_search"] = " ".join(str(v or "") for row in case["_findings"] for k, v in row.items() if not k.startswith("_")).casefold()
            self.refresh_case(case)
        previous = self.source
        self.source = FindingsModel(cases, self)
        self.proxy.setSourceModel(self.source)
        previous.deleteLater()
        self.restoring = True
        for combo, values in ((self.kind, ["Todos"] + sorted({str(row.get("Tipo_hallazgo")) for row in self.rows})), (self.ipress, ["Todas"] + sorted({c["_ipress"] for c in cases}))):
            combo.clear()
            combo.addItems(values)
        session = self.store.load_session() if self.store else {}
        self.query.setText(session.get("query", ""))
        for name in ("priority", "kind", "state", "ipress"):
            combo = getattr(self, name)
            value = session.get(name, "Pendientes" if name == "state" else "Todas" if name != "kind" else "Todos")
            index = combo.findText(value)
            combo.setCurrentIndex(max(0, index))
        self.restoring = False
        self.table.resizeColumnsToContents()
        self.table.setColumnHidden(3, True)
        self.table.setColumnHidden(4, True)
        self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Stretch)
        self.table.sortByColumn(0, Qt.SortOrder.AscendingOrder)
        self.apply_filters()
        self.select_case(session.get("case"))
        if self.current_case and session.get("page") and session.get("case") == self.current_case["_key"]:
            for index, row in enumerate(self.current_case["_findings"]):
                if row["_id"] == session.get("field"):
                    self.findings.selectRow(index)
                    break
            self.viewer.zoom.setCurrentIndex(session.get("zoom", 0))
            self.viewer.page.setValue(session["page"])
        if self.rows and self.rows[0].get("_warning"):
            self.feedback.setText(self.rows[0]["_warning"])
        self.update_actions()

    def refresh_case(self, case):
        decisions = self.store.decisions if self.store else {}
        choices = [decisions.get(row["_id"], {}).get("choice") for row in case["_findings"]]
        case["Estado"] = "Completado" if choices and all(c in RESOLVED for c in choices) else "Derivado" if choices and all(choices) and "No concluyente" in choices else "En revisión" if any(choices) else "Pendiente"

    def select_case(self, key):
        for i in range(self.proxy.rowCount()):
            case = self.source.rows[self.proxy.mapToSource(self.proxy.index(i, 0)).row()]
            if case["_key"] == key:
                self.table.selectRow(i)
                return
        if self.proxy.rowCount():
            self.table.selectRow(0)
        else:
            self.selected(QModelIndex(), QModelIndex())

    def apply_filters(self, *_args):
        if self.restoring:
            return
        self.proxy.beginFilterChange()
        for name in ("priority", "kind", "state", "ipress"):
            setattr(self.proxy, name, getattr(self, name).currentText())
        self.proxy.query = self.query.text().strip()
        self.proxy.endFilterChange(QSortFilterProxyModel.Direction.Rows)
        if not self.proxy.rowCount():
            self.selected(QModelIndex(), QModelIndex())
        elif not self.table.currentIndex().isValid():
            self.table.selectRow(0)
        self.update_count()
        self.save_session()

    def update_count(self):
        cases = self.source.rows
        complete = sum(c["Estado"] == "Completado" for c in cases)
        derived = sum(c["Estado"] == "Derivado" for c in cases)
        self.count.setText(f"{len(cases)} casos · {len(self.rows)} hallazgos\n{complete} completados · {derived} derivados · {len(cases)-complete-derived} pendientes · {self.proxy.rowCount()} visibles")

    def selected(self, index, _previous):
        if not index.isValid():
            self.current_case = None
            self.set_fields([])
            self.details.clear()
            self.viewer.open_pdf(None)
            self.update_actions()
            return
        self.current_case = self.source.rows[self.proxy.mapToSource(index).row()]
        row = self.current_case["_findings"][0]
        path = row.get("_pdf_path")
        if not path or not Path(path).is_file():
            relative = str(row.get("Ruta_relativa") or "").replace("\\", "/")
            if self.pdf_root and relative and not Path(relative).is_absolute():
                root = Path(self.pdf_root).resolve()
                candidate = (root / relative).resolve()
                path = str(candidate) if candidate.is_relative_to(root) and candidate.is_file() else None
        self.viewer.open_pdf(path)
        self.set_fields(self.current_case["_findings"])
        self.save_session()
        self.update_actions()

    def set_fields(self, rows):
        previous = self.findings_source
        decisions = self.store.decisions if self.store else {}
        self.findings_source = FindingsModel([{"Examen": str(row.get("Examen") or "Registro").replace("_", " "), "Hito 2": row.get("Valor_LLM_OCR") if row.get("Valor_LLM_OCR") is not None else "Sin dato", "Hito 3": row.get("Valor_SAIRC") if row.get("Valor_SAIRC") is not None else "Sin dato", "Decisión": decisions.get(row["_id"], {}).get("choice", "Pendiente")} for row in rows], self, ["Examen", "Hito 2", "Hito 3", "Decisión"])
        self.findings.setModel(self.findings_source)
        self.findings.selectionModel().currentRowChanged.connect(self.field_selected)
        previous.deleteLater()
        if rows:
            self.findings.selectRow(0)

    def active_row(self):
        index = self.findings.currentIndex().row()
        return self.current_case["_findings"][index] if self.current_case and 0 <= index < len(self.current_case["_findings"]) else None

    def field_selected(self, *_args):
        row = self.active_row()
        self.decision.setCurrentIndex(0)
        self.manual_value.clear()
        self.reason.clear()
        if row:
            if row.get("_warning"):
                self.feedback.setText(row["_warning"])
            self.details.setHtml("<b>" + escape(str(row.get("Examen") or "Registro")) + "</b><br>" + escape(str(row.get("Interpretacion_control") or row.get("Tipo_hallazgo") or "")))
            if self.store:
                saved = self.store.decisions.get(row["_id"], {})
                self.reason.setText(saved.get("reason", ""))
                if saved:
                    self.details.append("<b>Decisión guardada:</b> " + escape(str(saved.get("choice"))) + " · Valor confirmado: " + escape(str(saved.get("value") if saved.get("value") is not None else "Sin dato")))
                    self.manual_value.setText(str(saved.get("value") if saved.get("value") is not None else ""))
                if saved.get("page") and self.viewer.path:
                    self.viewer.page.setValue(saved["page"])
        self.update_actions()

    def save_decision(self, advance=False):
        row = self.active_row()
        if not row or not self.store:
            return
        try:
            self.store.save(row["_id"], self.decision.currentText(), self.manual_value.text(), self.reason.text(), self.viewer.page.value() if self.viewer.path else None)
            self.refresh_case(self.current_case)
            index = self.findings.currentIndex().row()
            self.findings_source.rows[index]["Decisión"] = self.store.decisions.get(row["_id"], {}).get("choice", "Pendiente")
            self.findings_source.dataChanged.emit(self.findings_source.index(index, 0), self.findings_source.index(index, 3))
            case_index = self.source.rows.index(self.current_case)
            self.source.dataChanged.emit(self.source.index(case_index, 0), self.source.index(case_index, 5))
            self.feedback.setText("Decisión guardada. Propuesta separada; SQL y fuentes permanecen intactos.")
            self.update_count()
            self.save_session()
            self.update_actions()
            if advance:
                undecided = [i for i, finding in enumerate(self.current_case["_findings"]) if finding["_id"] not in self.store.decisions]
                if undecided:
                    self.findings.selectRow(undecided[0])
                else:
                    self.next_pending()
        except Exception as exc:
            self.feedback.setText(f"No se guardó: {exc}")

    def next_pending(self):
        start = self.table.currentIndex().row()
        for step in range(1, self.proxy.rowCount() + 1):
            i = (start + step) % self.proxy.rowCount()
            case = self.source.rows[self.proxy.mapToSource(self.proxy.index(i, 0)).row()]
            if case["Estado"] in {"Pendiente", "En revisión"}:
                self.table.selectRow(i)
                self.apply_filters()
                return
        self.apply_filters()
        self.feedback.setText("No quedan pendientes en los filtros actuales. Los derivados requieren resolución aparte.")

    def step(self, delta):
        if self.proxy.rowCount():
            row = max(0, min(self.proxy.rowCount() - 1, self.table.currentIndex().row() + delta))
            self.table.selectRow(row)

    def undo(self):
        if not self.store:
            return
        try:
            key = self.store.undo()
            if not key:
                self.feedback.setText("No hay decisiones para deshacer en esta versión de fuentes.")
                return
            for case in self.source.rows:
                self.refresh_case(case)
            self.source.dataChanged.emit(self.source.index(0, 0), self.source.index(len(self.source.rows)-1, 5))
            self.apply_filters()
            self.select_case(case_key(self.store.rows[key]))
            self.feedback.setText("Decisión deshecha; historial conservado.")
        except Exception as exc:
            self.feedback.setText(f"No se pudo deshacer: {exc}")

    def update_actions(self):
        enabled = bool(self.store and self.active_row())
        self.save_button.setEnabled(enabled)
        self.next_button.setEnabled(enabled)
        self.decision.setEnabled(enabled)
        self.undo_button.setEnabled(bool(self.store))
        self.export_button.setEnabled(bool(self.store))

    def save_session(self):
        if self.restoring or not self.store:
            return
        payload = {name: getattr(self, name).currentText() for name in ("priority", "kind", "state", "ipress")}
        row = self.active_row()
        payload.update(query=self.query.text(), case=self.current_case["_key"] if self.current_case else None, field=row["_id"] if row else None, page=self.viewer.page.value(), zoom=self.viewer.zoom.currentIndex())
        try:
            self.store.save_session(payload)
        except Exception as exc:
            self.feedback.setText(f"No se guardó la posición: {exc}")

    def filter_dialog(self):
        dialog = QDialog(self)
        dialog.setWindowTitle("Filtros de revisión")
        layout = QFormLayout(dialog)
        layout.addRow("IPRESS", self.ipress)
        layout.addRow("Tipo de hallazgo", self.kind)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
        buttons.rejected.connect(dialog.accept)
        layout.addRow(buttons)
        dialog.exec()
        self.ipress.setParent(self)
        self.kind.setParent(self)
        self.ipress.hide()
        self.kind.hide()

    def show_details(self):
        row = self.active_row()
        if not row:
            return
        dialog = QDialog(self)
        dialog.setWindowTitle("Fuentes, correlación y cobertura")
        dialog.resize(650, 550)
        layout = QVBoxLayout(dialog)
        detail = QTextBrowser()
        detail.setHtml("<p>Comparación cuantitativa: " + escape(", ".join(row.get("_coverage") or []) or "Reporte anterior sin cobertura registrada") + ". Serologías y campos fuera del mapeo no están comparados. Prioridad de revisión; no gravedad clínica.</p>" + "<br>".join(f"<b>{escape(k.replace('_',' '))}:</b> {escape(str(v if v is not None else 'Sin dato'))}" for k,v in row.items() if not k.startswith("_")))
        context = row.get("_context", [])
        if context:
            detail.append("<h3>Valores comparados de este PDF</h3><table><tr><th>Examen</th><th>Hito 2</th><th>Hito 3</th><th>Estado</th><th>Correlación</th></tr>" + "".join("<tr>" + "".join("<td>" + escape(str(item.get(key) if item.get(key) is not None else "Sin dato")) + "</td>" for key in ("Examen", "Hito_2", "Hito_3", "Estado", "Correlacion")) + "</tr>" for item in context) + "</table>")
        layout.addWidget(detail)
        dialog.exec()

    def export_decisions(self):
        if not self.store:
            return
        target, _ = QFileDialog.getSaveFileName(self, "Exportar propuestas, sin aplicar a fuentes", "propuestas_revision.xlsx", "Excel (*.xlsx)")
        if target:
            try:
                self.store.export(target)
                self.feedback.setText("Propuestas exportadas. No aplicadas a SQL ni a los archivos originales.")
            except Exception as exc:
                self.feedback.setText(f"No se pudo exportar: {exc}")

    def finish(self):
        self.save_session()
        if self.store:
            self.store.close()
            self.store = None
        self.viewer.document.close()

    def closeEvent(self, event):
        self.finish()
        event.accept()
